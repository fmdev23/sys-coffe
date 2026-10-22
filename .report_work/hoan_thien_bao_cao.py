import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
import json, re, ast, textwrap
from collections import defaultdict
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

W=Path(__file__).parent
ROOT=W.parent
OUT=ROOT/'docs/BaoCao_QuanLyQuanCaPhe.docx'
old=json.loads((W/'old_content.json').read_text(encoding='utf-8'))
evidence=json.loads((W/'evidence_new.json').read_text(encoding='utf-8'))
checklist=re.findall(r'^\d+\. (.+)$',(ROOT/'README.md').read_text(encoding='utf-8').split('## Chạy')[0],re.M)
if len(checklist)!=67:
    checklist=json.loads((W/'checklist.json').read_text(encoding='utf-8'))
else:
    (W/'checklist.json').write_text(json.dumps(checklist,ensure_ascii=False),encoding='utf-8')
doc=Document()
styles=doc.styles
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','TOC Heading']:
    st=styles[name]
    st.font.name='Times New Roman'
    st.font.size=Pt(13)
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.2
    st.paragraph_format.space_before=Pt(6)
    st.paragraph_format.space_after=Pt(0)
styles['Normal'].paragraph_format.first_line_indent=Cm(1)
styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3']:
    styles[name].font.bold=True
    styles[name].paragraph_format.keep_with_next=True
styles['Heading 1'].font.size=Pt(14)
styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 1'].paragraph_format.first_line_indent=Cm(0)
styles['Heading 1'].paragraph_format.page_break_before=True
styles['Heading 2'].paragraph_format.first_line_indent=Cm(0)
styles['Heading 3'].paragraph_format.first_line_indent=Cm(1)
for name,outline in [('Front Heading',0),('Appendix Heading',9),('Code',9),('References',9)]:
    st=styles.add_style(name,1)
    st.base_style=styles['Normal']
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    level=OxmlElement('w:outlineLvl')
    level.set(qn('w:val'),str(outline))
    st.element.get_or_add_pPr().append(level)
styles['Front Heading'].font.size=Pt(14)
styles['Front Heading'].font.bold=True
styles['Front Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Front Heading'].paragraph_format.page_break_before=True
styles['Front Heading'].paragraph_format.keep_with_next=True
styles['Appendix Heading'].font.bold=True
styles['Appendix Heading'].paragraph_format.keep_with_next=True
styles['Code'].paragraph_format.line_spacing=1.0
styles['Code'].paragraph_format.space_before=Pt(0)
styles['Code'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
styles['References'].paragraph_format.line_spacing=1.0
styles['References'].paragraph_format.left_indent=Cm(1)
styles['References'].paragraph_format.first_line_indent=Cm(-1)
styles['References'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
styles['Caption'].font.bold=False
styles['Caption'].font.italic=False
styles['Caption'].paragraph_format.first_line_indent=Cm(0)
styles['Caption'].paragraph_format.line_spacing=1.0
for lev in range(1,4):
    name='TOC '+str(lev)
    st=styles[name] if name in styles else styles.add_style(name,1)
    st.base_style=styles['Normal']
    st.font.name='Times New Roman'
    st.font.size=Pt(13)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.left_indent=Cm((lev-1)*.5)
    st.paragraph_format.line_spacing=1.0
    st.paragraph_format.space_before=Pt(3)
    st.paragraph_format.space_after=Pt(0)
for st in styles:
    if st.type==1:
        st.font.name='Times New Roman'
        fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+attr),None)

def page_setup(sec):
    sec.page_width=Cm(21)
    sec.page_height=Cm(29.7)
    sec.left_margin=Cm(3.5)
    sec.right_margin=Cm(2)
    sec.top_margin=Cm(2)
    sec.bottom_margin=Cm(2)
    sec.header_distance=Cm(1)
    sec.footer_distance=Cm(1)

def field(pp,code,value=''):
    run=pp.add_run()
    begin=OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText')
    instr.set(qn('xml:space'),'preserve')
    instr.text=' '+code+' '
    separate=OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'),'separate')
    t=OxmlElement('w:t')
    t.text=str(value)
    end=OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'),'end')
    for item in [begin,instr,separate,t,end]:run._r.append(item)

def set_pages(sec,fmt):
    sec.footer.is_linked_to_previous=False
    pp=sec.footer.paragraphs[0]
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.line_spacing=1
    field(pp,'PAGE','1')
    el=OxmlElement('w:pgNumType')
    el.set(qn('w:fmt'),fmt)
    el.set(qn('w:start'),'1')
    sec._sectPr.append(el)

# Đánh số mục bằng danh sách nhiều cấp của Word.
numbering=doc.part.numbering_part.element
for abstract_id,num_id,formats in [
    (41,41,[(0,'CHƯƠNG %1'),(1,'%1.%2'),(2,'%1.%2.%3')]),
    (42,42,[(0,'%1')])]:
    abstract=OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'),str(abstract_id))
    for tag,val in [('nsid',f'15BA{abstract_id:04X}'),('multiLevelType','multilevel'),('tmpl',f'15BA{abstract_id:04X}')]:
        el=OxmlElement('w:'+tag)
        el.set(qn('w:val'),val)
        abstract.append(el)
    for lev,fmt in formats:
        lvl=OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'),str(lev))
        for tag,val in [('start','1'),('numFmt','decimal'),('lvlText',fmt),('suff','space')]:
            el=OxmlElement('w:'+tag)
            el.set(qn('w:val'),val)
            lvl.append(el)
        ppr=OxmlElement('w:pPr')
        ind=OxmlElement('w:ind')
        ind.set(qn('w:left'),'0')
        ind.set(qn('w:hanging'),'0')
        ppr.append(ind)
        lvl.append(ppr)
        abstract.append(lvl)
    first_num=numbering.find(qn('w:num'))
    numbering.insert(list(numbering).index(first_num),abstract)
    num=OxmlElement('w:num')
    num.set(qn('w:numId'),str(num_id))
    ref=OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'),str(abstract_id))
    num.append(ref)
    numbering.append(num)

def num_para(pp,level,ident=41):
    num=OxmlElement('w:numPr')
    for tag,val in [('ilvl',level),('numId',ident)]:
        e=OxmlElement('w:'+tag)
        e.set(qn('w:val'),str(val))
        num.append(e)
    pp._p.get_or_add_pPr().append(num)

chapter=''
counts=defaultdict(int)
figures=[]
tables=[]

def p(text,style=None):
    assert ';' not in text, text
    pp=doc.add_paragraph(style=style)
    for chunk in re.split(r'(Data Flow Diagram|Comma-Separated Values|Open Database Connectivity|Structured Query Language|Password-Based Key Derivation Function 2|Unicode Transformation Format|Byte Order Mark)',text):
        rr=pp.add_run(chunk)
        rr.italic=chunk in ('Data Flow Diagram','Comma-Separated Values','Open Database Connectivity','Structured Query Language','Password-Based Key Derivation Function 2','Unicode Transformation Format','Byte Order Mark')
    return pp

def h(title,lev=2):
    pp=doc.add_heading(title.upper() if lev==2 else title,lev)
    num_para(pp,lev-1)
    if lev==3:
        pp.paragraph_format.left_indent=Cm(1)
        pp.paragraph_format.first_line_indent=Cm(0)
    if title in ['Thiết kế dữ liệu','Thuật toán lưu đơn hàng','Thuật toán thanh toán']:
        pp.paragraph_format.page_break_before=True
    return pp

def ch(no,title):
    global chapter
    chapter=str(no)
    pp=doc.add_heading('\n'+title.upper(),1)
    num_para(pp,0)

def front(title):
    return p(title.upper(),'Front Heading')

def cap(label,title):
    key=(label,chapter)
    counts[key]+=1
    no=(chapter+'.' if chapter else '')+str(counts[key])
    pp=doc.add_paragraph(style='Caption')
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if label=='Hình' else WD_ALIGN_PARAGRAPH.LEFT
    pp.paragraph_format.keep_with_next=label=='Bảng'
    pp.add_run(label+' '+(chapter+'.' if chapter else ''))
    field(pp,('SEQ Figure' if label=='Hình' else 'SEQ Table')+(r' \r 1' if counts[key]==1 else '')+r' \* ARABIC',counts[key])
    pp.add_run(' '+title)
    (figures if label=='Hình' else tables).append({'number':no,'title':title,'chapter':chapter})
    return pp

def figure(file,title,maxh=17.5,width=15.5,break_before=False):
    number=(chapter+'.' if chapter else '')+str(counts[('Hình',chapter)]+1)
    ref=p('Hình '+number+' trình bày '+title[0].lower()+title[1:]+'.')
    ref.paragraph_format.keep_with_next=True
    if break_before:ref.paragraph_format.page_break_before=True
    path=ROOT/file
    img=Image.open(path)
    width=min(width,maxh*img.width/img.height)
    pp=doc.add_paragraph()
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.line_spacing=1.0
    pp.paragraph_format.keep_with_next=True
    pp.add_run().add_picture(str(path),width=Cm(width))
    cap('Hình',title)

def table(title,rows,widths=None):
    number=(chapter+'.' if chapter else '')+str(counts[('Bảng',chapter)]+1)
    ref=p('Bảng '+number+' tổng hợp '+title[0].lower()+title[1:]+'.')
    ref.paragraph_format.keep_with_next=True
    cap('Bảng',title)
    tt=doc.add_table(rows=0,cols=len(rows[0]))
    tt.style='Table Grid'
    tt.autofit=False
    if widths is None:widths=[15.5/len(rows[0])]*len(rows[0])
    for col,w in zip(tt.columns,widths):col.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','000000')]:b.set(qn('w:'+k),v)
        borders.append(b)
    tt._tbl.tblPr.append(borders)
    for i,row in enumerate(rows):
        rr=tt.add_row()
        cant=OxmlElement('w:cantSplit')
        rr._tr.get_or_add_trPr().append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader')
            rr._tr.trPr.append(repeat)
        for cell,txt,w in zip(rr.cells,row,widths):
            cell.width=Cm(w)
            cp=cell.paragraphs[0]
            cp.paragraph_format.first_line_indent=Cm(0)
            cp.paragraph_format.line_spacing=1.0
            cp.paragraph_format.space_before=Pt(0)
            cp.paragraph_format.space_after=Pt(0)
            cp.paragraph_format.keep_with_next=i<len(rows)-1
            cp.alignment=WD_ALIGN_PARAGRAPH.RIGHT if i and re.fullmatch(r'[0-9 .,]+',str(txt)) else WD_ALIGN_PARAGRAPH.LEFT
            run=cp.add_run(str(txt).replace(' × ',' nhân '))
            run.bold=i==0
            run.font.size=Pt(13)
            if i==0:
                shade=OxmlElement('w:shd')
                shade.set(qn('w:fill'),'F2F2F2')
                cell._tc.get_or_add_tcPr().append(shade)
    return tt

def extract(file,name):
    source=(ROOT/file).read_text(encoding='utf-8')
    node=next(n for n in ast.walk(ast.parse(source)) if isinstance(n,ast.FunctionDef) and n.name==name)
    return textwrap.dedent('\n'.join(source.splitlines()[node.lineno-1:node.end_lineno]))

def code(text):
    pp=p(text,'Code')
    pp.paragraph_format.keep_together=True
    return pp

sec=doc.sections[0]
page_setup(sec)
def cover(text,size=13,bold=False,color=None,after=9):
    pp=p(text)
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.line_spacing=1.15
    pp.paragraph_format.space_before=Pt(0)
    pp.paragraph_format.space_after=Pt(after)
    for rr in pp.runs:
        rr.font.size=Pt(size)
        rr.bold=bold
        if color:rr.font.color.rgb=RGBColor.from_string(color)
    return pp
cover('TRƯỜNG ĐẠI HỌC NAM CẦN THƠ',14,True)
cover('TRƯỜNG CÔNG NGHỆ SỐ VÀ TRÍ TUỆ NHÂN TẠO',13,True,after=15)
pp=cover('',after=14)
pp.add_run().add_picture(str(ROOT/'sodo/logo_truong.png'),height=Cm(2.9))
cover('BÁO CÁO HỌC PHẦN',20,True,'17365D',15)
cover('LẬP TRÌNH PYTHON - THỰC HÀNH',17,True,'C00000',22)
cover('Đề tài',13,False,after=8)
cover('XÂY DỰNG ỨNG DỤNG\nQUẢN LÝ QUÁN CÀ PHÊ BẰNG PYTHON',17,True,after=13)
pp=cover('Coffee Shop Management Application',14,after=25)
pp.runs[0].italic=True
cover('Sinh viên thực hiện',13,True,after=6)
cover('Lê Minh Phú - MSSV: 236717\nLớp: DH23TIN02',13,after=15)
cover('Giảng viên hướng dẫn\nBùi Thị Diễm Trinh',13,after=23)
cover('Cần Thơ, tháng 10 năm 2026',13,after=0)

sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
page_setup(sec)
set_pages(sec,'lowerRoman')
for title,instruction in [('MỤC LỤC',r'TOC \o "1-3" \h \z \u'),('DANH MỤC BẢNG',r'TOC \h \z \c "Table"'),('DANH MỤC HÌNH',r'TOC \h \z \c "Figure"')]:
    pp=p(title)
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.page_break_before=title!='MỤC LỤC'
    pp.paragraph_format.keep_with_next=True
    pp.runs[0].bold=True
    pp.runs[0].font.size=Pt(14)
    field(doc.add_paragraph(),instruction,'')
pp=p('DANH MỤC TỪ VIẾT TẮT')
pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp.paragraph_format.page_break_before=True
pp.paragraph_format.first_line_indent=Cm(0)
pp.runs[0].bold=True
pp.runs[0].font.size=Pt(14)
abbreviations=[('CSV','Comma-Separated Values, tệp có các giá trị phân cách bằng dấu phẩy'),('DFD','Data Flow Diagram, sơ đồ luồng dữ liệu'),('ODBC','Open Database Connectivity, giao tiếp kết nối cơ sở dữ liệu'),('PBKDF2','Password-Based Key Derivation Function 2, hàm dẫn xuất khóa từ mật khẩu'),('SQL','Structured Query Language, ngôn ngữ truy vấn có cấu trúc'),('UTF-8','Unicode Transformation Format 8-bit, mã hóa Unicode')]
for short,meaning in abbreviations:
    pp=p(short+': '+meaning)
    pp.paragraph_format.first_line_indent=Cm(0)

sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
page_setup(sec)
set_pages(sec,'decimal')
pp=front('Giới thiệu')
pp.paragraph_format.page_break_before=False
for title,body in [
 ('Đặt vấn đề',[
 'Quán cà phê cần theo dõi bàn đang phục vụ, món đã gọi và tiền phải thu trong từng đơn. Khi ghi nhận rời rạc, người phục vụ khó kiểm tra món đã lưu và quản lý mất thời gian tổng hợp doanh thu. Đề tài xây dựng một ứng dụng Python trên máy tính để thực hiện các thao tác này trong cùng một quy trình.',
 'Ứng dụng sử dụng Tkinter/ttk để tạo giao diện tiếng Việt và Microsoft SQL Server để lưu dữ liệu. SQL là ngôn ngữ truy vấn có cấu trúc (Structured Query Language). Chương trình kết nối qua giao tiếp ODBC (Open Database Connectivity) bằng thư viện pyodbc. Nội dung báo cáo mô tả đúng chức năng đang có và kết quả chạy thử của dự án.']),
 ('Những nghiên cứu liên quan',[])]:
    pp=doc.add_heading(title,2)
    num_para(pp,0,42)
    for text in body:p(text)
table('Ứng dụng và thư viện được khảo sát',old['tables'][0],[3.7,5.5,6.3])
pp=p('Nguồn: tổng hợp từ abdelbend (không ghi năm a, b), mkleehammer và cộng sự (không ghi năm b), Python Software Foundation (không ghi năm a, b, d, e).')
pp.paragraph_format.first_line_indent=Cm(0)
pp.paragraph_format.line_spacing=1.0
pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
for rr in pp.runs:
    rr.italic=True
    rr.font.size=Pt(11)
p('Restaurant Management System là ứng dụng tham khảo chính, có mã nguồn công khai bằng Python (abdelbend, không ghi năm a, b). Các thư viện còn lại có tài liệu chính thức: Tkinter cho giao diện, csv cho tập tin, decimal cho số tiền và unittest cho kiểm thử (Python Software Foundation, không ghi năm a, b, d, e). pyodbc cung cấp kết nối cơ sở dữ liệu (mkleehammer và cộng sự, không ghi năm b).')
p('Ứng dụng tham khảo tổ chức việc chọn món, ghi nhận yêu cầu và hiển thị hóa đơn. Mã db.py sử dụng SQLite với dữ liệu thực đơn và khách hàng. Cách đặt thao tác chọn món cạnh thông tin đơn phù hợp một chương trình phục vụ tại quầy. Tuy nhiên, cấu trúc lưu trữ và giao diện của ứng dụng đó không được xem là bằng chứng cho chức năng của đề tài này. Các chức năng phân quyền, bàn, điểm khách và báo cáo trong đề tài được đối chiếu trực tiếp với mã nguồn hiện có.')
p('Đề tài áp dụng cách tổ chức thao tác ngắn, hiển thị tổng tiền gần giỏ hàng và giữ trạng thái bàn rõ ràng. Dữ liệu được lưu ở SQL Server với đơn và chi tiết đơn tách riêng. Việc phân tích ứng dụng tham khảo giúp xác định cách trình bày quy trình, còn các quyết định xử lý được kiểm tra bằng mã và kết quả chạy của dự án.')
for title,body in [
 ('Mục tiêu',[
 'Mục tiêu là xây dựng ứng dụng quản lý quán cà phê quy mô nhỏ và vừa, có giao diện dễ thao tác, lưu được dữ liệu và hỗ trợ theo dõi doanh thu. Chương trình cần phân biệt quyền quản lý và nhân viên, kiểm tra dữ liệu nhập, giữ đúng lịch sử giá bán và xử lý trọn vẹn việc thanh toán.',
 'Về học phần, dự án thể hiện đủ tám nội dung đã học thông qua biến, điều kiện, vòng lặp, hàm, chuỗi, danh sách, tập tin và giao diện Tkinter. Mã được chia theo nhiệm vụ, sử dụng các câu lệnh quen thuộc để có thể trình bày và kiểm tra từng bước.']),
 ('Đối tượng và phạm vi nghiên cứu',[
 'Đối tượng sử dụng là quản lý và nhân viên phục vụ tại quầy. Quản lý chuẩn bị danh mục, nhân sự, tài khoản và xem báo cáo. Nhân viên xem thực đơn, theo dõi bàn, lập hoặc sửa đơn và xác nhận thanh toán. Khách hàng được lưu như dữ liệu của đơn, chưa có giao diện đăng nhập riêng.',
 'Phạm vi triển khai là ứng dụng trên Windows, kết nối SQL Server bằng xác thực Windows. Chương trình chưa có quản lý kho nguyên liệu, thuế, khuyến mãi, in hóa đơn, thanh toán ngân hàng hoặc vận hành qua Internet. Chức năng tập tin hiện có là xuất báo cáo CSV (Comma-Separated Values), không có nhập dữ liệu từ CSV.']),
 ('Sơ đồ DFD cấp 0',[
 'Sơ đồ luồng dữ liệu (Data Flow Diagram, DFD) cấp 0 biểu diễn toàn bộ ứng dụng bằng một tiến trình trung tâm. Hình 1 sử dụng nguyên bản bản vẽ do sinh viên cung cấp, có hai tác nhân Nhân viên và Quản lý. Tiến trình được đánh số 1 theo bản gốc. Luồng vào gồm thông tin đăng nhập, dữ liệu đơn, danh mục và yêu cầu thống kê. Luồng ra gồm kết quả xác thực, thông tin phục vụ và báo cáo.'])]:
    pp=doc.add_heading(title,2)
    num_para(pp,0,42)
    for text in body:p(text)
figure('docs/DFD_Muc0.png','Sơ đồ DFD cấp 0 của ứng dụng',width=15.5)
p('Ảnh DFD gốc có kích thước 927 nhân 180 điểm ảnh nên chữ trong hình nhỏ khi đặt trên trang A4. Bảng 2 chép lại đúng các nhãn và chiều truyền để thuận tiện đọc và đối chiếu, không chỉnh sửa ảnh gốc.')
table('Các luồng dữ liệu trên bản DFD gốc',[
 ['Nguồn','Đích','Dữ liệu trao đổi'],
 ['Nhân viên','Hệ thống','yêu cầu sửa/hủy đơn'],
 ['Nhân viên','Hệ thống','thông tin bàn, món, số lượng, tiền khách đưa'],
 ['Nhân viên','Hệ thống','thông tin đăng nhập'],
 ['Hệ thống','Nhân viên','kết quả xác thực'],
 ['Hệ thống','Nhân viên','chi tiết đơn, tổng tiền, tiền thừa'],
 ['Hệ thống','Nhân viên','thực đơn, trạng thái bàn'],
 ['Quản lý','Hệ thống','thông tin đăng nhập'],
 ['Quản lý','Hệ thống','dữ liệu sản phẩm, bàn, khách hàng, nhân viên, tài khoản'],
 ['Quản lý','Hệ thống','yêu cầu thu tiền, khoảng thời gian thống kê'],
 ['Hệ thống','Quản lý','doanh thu, số hóa đơn, món bán chạy, báo cáo CSV'],
 ['Hệ thống','Quản lý','danh sách sản phẩm, bàn, khách hàng, nhân viên, tài khoản'],
 ['Hệ thống','Quản lý','kết quả xác thực']], [2.8,2.8,9.9])
p('Khi quản lý chọn xuất báo cáo, thao tác từ giao diện được xử lý bởi StatisticsTab.export() trong Python, đọc số liệu qua CoffeeService.statistics() rồi ghi tập tin CSV ở vị trí người dùng chọn. SQL Server là nơi lưu dữ liệu chính. Tập tin CSV là đầu ra tổng hợp, không thay thế cơ sở dữ liệu và không được nhập ngược tự động.')
p('Bản vẽ phù hợp mức mô tả tổng quát. Luồng “yêu cầu thu tiền, khoảng thời gian thống kê” của Quản lý đang gộp hai nhóm dữ liệu, có thể tách khi cần diễn giải chi tiết. Kết quả lưu, sửa, hủy và thông báo lỗi chưa được ghi thành nhãn riêng. Các nhận xét này không làm thay đổi bản gốc. Mã PlantUML tương đương và kết quả đối chiếu được lưu trong README.md để sinh viên kiểm chứng.')

ch(1,'Đặc tả yêu cầu')
h('Yêu cầu chức năng')
table('Các chức năng của chương trình',[row[1:] for row in old['tables'][1]],[4.5,11])
p('Luồng phục vụ bắt đầu từ việc chọn bàn hoặc mở đơn đang xử lý. Người dùng thêm món, điều chỉnh số lượng, chọn khách nếu có và lưu đơn. Sau khi nhận tiền, người dùng xác nhận thanh toán. Chương trình ghi trạng thái đơn, thời điểm thanh toán, điểm tích lũy và trạng thái bàn trong cùng một giao dịch.')
h('Quy tắc dữ liệu và nghiệp vụ')
table('Quy tắc kiểm tra dữ liệu',old['tables'][2],[4.2,11.3])
p('Tổng đơn được cộng từ thành tiền của từng món. Ví dụ, hai cà phê đen giá 25.000 đồng và một cà phê sữa giá 30.000 đồng tạo tổng 80.000 đồng. Nếu nhận 100.000 đồng, giao diện hiển thị tiền thừa 20.000 đồng. Khách đăng ký nhận 8 điểm khi đơn được thanh toán. Đơn của khách lẻ không cộng điểm.')
h('Phân quyền sử dụng')
table('Quyền thao tác theo vai trò',old['tables'][3],[7.5,3,5])
p('Quản lý có tám trang chức năng. Nhân viên có Đơn hàng, Sơ đồ bàn và Thực đơn. Các hàm nghiệp vụ kiểm tra lại phiên đăng nhập và vai trò trước khi thao tác, vì việc ẩn nút trên giao diện chưa đủ để bảo vệ dữ liệu.')
h('Yêu cầu giao diện và chất lượng')
p('Các thao tác chính dùng màu xanh đậm, nền sáng và khoảng cách thống nhất. Màu trạng thái luôn đi cùng nhãn chữ. Bàn được biểu diễn bằng thẻ để người phục vụ nhận ra tình trạng nhanh. Màn hình gọi món đặt thực đơn bên trái và giỏ hàng bên phải. Nút lưu và thu tiền nằm gần tổng đơn để giảm chuyển hướng khi thao tác.')
p('Nút, ô nhập và danh sách chọn vẫn là thành phần ttk. ui/theme.py thay hình nền và trạng thái hiển thị để có góc bo, phản hồi khi rê chuột, nhận tiêu điểm hoặc bị vô hiệu hóa. Các ảnh nền được Tkinter đọc trực tiếp, không cần bổ sung thư viện giao diện. Dữ liệu liên kết và giao dịch được kiểm tra để hạn chế lưu dở dang hoặc làm mất lịch sử.')

ch(2,'Thiết kế giải pháp')
h('Sơ đồ phân cấp hàm')
p('Các nhóm chức năng được liên hệ với tên hàm thật trong mã nguồn. login() là hàm ở models/service.py. show_login() thuộc CoffeeApp, các hàm xử lý dữ liệu thuộc CoffeeService. Nhóm báo cáo sử dụng Dashboard.refresh() và StatisticsTab.export(). Đường nối trong sơ đồ thể hiện phân rã chức năng, không biểu diễn thứ tự thực thi.')
figure('sodo/phan_cap_ham.png','Sơ đồ phân cấp hàm',maxh=12)
h('Tổ chức mã nguồn')
roles=[row[:] for row in old['tables'][4]]
roles[3][1]='Mở kết nối SQL Server và chuyển kết quả truy vấn thành danh sách từ điển.'
table('Vai trò các thành phần mã nguồn',roles,[5,10.5])
p('Giao diện nhận dữ liệu và gọi hàm xử lý. CoffeeService kiểm tra quyền, trạng thái và thực hiện truy vấn. Những thao tác ghi dùng khối try/except/finally với commit(), rollback() và close() rõ ràng. Cách tổ chức này giữ được ranh giới giữa hiển thị và xử lý mà không bổ sung framework hay lớp trung gian không cần thiết.')
h('Thiết kế dữ liệu')
p('Dự án có bảy bảng SQL. Sơ đồ quan hệ được bổ sung để làm rõ dữ liệu lưu trữ thực tế, phục vụ giải thích kết nối cơ sở dữ liệu. Tên bảng, trường và quan hệ lấy từ database_setup.py. Hai vạch biểu thị đúng một, vòng tròn kèm vạch biểu thị không hoặc một, vòng tròn kèm chân quạ biểu thị không hoặc nhiều.')
figure('sodo/quan_he_du_lieu.png','Quan hệ giữa bảy bảng dữ liệu',maxh=17.0)
p('Mỗi tài khoản bắt buộc liên kết một nhân viên, mỗi nhân viên có tối đa một tài khoản. Đơn bắt buộc có bàn và nhân viên lập, khách hàng có thể để trống. Một đơn đã lưu qua hàm nghiệp vụ phải có ít nhất một món. Riêng ràng buộc khóa ngoại không bắt buộc đơn luôn có chi tiết, vì vậy phía chi tiết trên sơ đồ thể hiện không hoặc nhiều.')
table('Các bảng và trường dữ liệu',old['tables'][5],[4.3,11.2])
p('Gia, DonGia và TongTien dùng DECIMAL(18,2), số lượng dùng số nguyên. ThanhTien là cột tính toán từ SoLuong và DonGia. PhienBan dùng ROWVERSION để phát hiện dữ liệu đã thay đổi kể từ lần đọc trước, không phải thời gian tạo hoặc sửa (Microsoft, không ghi năm). Các khóa ngoại giữ liên kết lịch sử. Chỉ mục duy nhất có điều kiện ngăn một bàn có hai đơn đang xử lý.')
h('Danh sách, từ điển và tập tin')
p('Danh sách (list) lưu nhiều phần tử theo thứ tự. Từ điển (dict) lưu cặp khóa và giá trị. Bộ giá trị (tuple) lưu một nhóm phần tử, được dùng cho các mốc thời gian và doanh thu trả về từ truy vấn.')
structures=[r[:] for r in old['tables'][6]]
for row,label in zip(structures[1:],['Danh sách từ điển','Giỏ hàng dạng từ điển','Ánh xạ lựa chọn','Gộp số lượng quantities','Danh sách bộ giá trị']):
    row[0]=label
table('Cấu trúc dữ liệu trong bộ nhớ',structures,[4.2,11.3])
p('Kết quả truy vấn được chuyển thành danh sách từ điển để truy cập bằng tên cột. Giỏ hàng dùng mã sản phẩm làm khóa nên thêm cùng món sẽ cập nhật số lượng. Trước khi ghi, quantities gộp các dòng cùng mã sản phẩm và kiểm tra giới hạn số lượng. Cấu trúc này được dùng trực tiếp trong vòng lặp của save_order().')
table('Cấu trúc tập tin CSV được xuất',old['tables'][7],[2.2,13.3])
p('StatisticsTab.export() ghi tập tin bằng csv.writer với encoding="utf-8-sig". Đây là mã hóa UTF-8 (Unicode Transformation Format 8-bit) có dấu nhận biết đầu tệp (Byte Order Mark). Tập tin dùng dấu phẩy phân cách các trường, ngày theo năm-tháng-ngày và số tiền dạng chuỗi thập phân của Decimal. Số liệu hiển thị trong báo cáo này dùng dấu chấm hàng nghìn theo quy định trình bày. Đó là cách trình bày khác với dữ liệu gốc trong tập tin.')
h('Thuật toán lưu đơn hàng')
p('save_order() kiểm tra giỏ trước khi mở giao dịch. Trong giao dịch, hàm kiểm tra phiên, đơn đang sửa, bàn và khách. Vòng lặp đọc sản phẩm, chọn đơn giá phù hợp và cộng tổng. Món đã lưu giữ tên và đơn giá cũ. Món ngừng bán chỉ được giữ số lượng không vượt chi tiết cũ. Sau khi ghi đơn và chi tiết, hàm đồng bộ bàn rồi xác nhận giao dịch.')
figure('sodo/luu_don_hang.png','Lưu đồ thuật toán lưu đơn hàng',maxh=16.8)
p('Nếu lỗi xảy ra trong giao dịch, khối except gọi rollback() rồi chuyển lỗi lên giao diện. Khối finally luôn đóng kết nối. Khi sửa đơn, expected_version phải khớp PhienBan hiện tại. Khi chuyển bàn, trạng thái cả bàn cũ và bàn mới đều được tính lại.')
h('Thuật toán thanh toán')
p('PaymentDialog kiểm tra tiền khách đưa và hiển thị tiền thừa. Khi người dùng xác nhận đã thu đủ, giao diện gọi finish_order() cùng phiên bản đơn. Hàm kiểm tra trạng thái, tính điểm, cập nhật đơn, khách và bàn trong cùng một giao dịch. Tiền nhận và tiền thừa chỉ hỗ trợ thao tác tại quầy, không có cột lưu hai giá trị này trong cơ sở dữ liệu.')
figure('sodo/thanh_toan_don.png','Lưu đồ thuật toán thanh toán đơn',maxh=16.8)
p('finish_order() còn có nhánh pay=False để hủy đơn đang xử lý. Nhánh đó ghi trạng thái “Đã huỷ” theo hằng CANCELLED trong mã, không ghi doanh thu và không cộng điểm. Lưu đồ trên trình bày nhánh thanh toán là chức năng xử lý quan trọng.')

ch(3,'Cài đặt giải pháp')
h('Môi trường và khởi động')
env=[r[:] for r in old['tables'][8]]
env[5][1]='Express 64-bit, phiên bản 16.0.1210.3.'
env[2][1]='3.14.3'
env[4][1]='5.3.0'
table('Môi trường chạy được kiểm tra ngày 20/10/2026',env,[4.3,11.2])
p('Môi trường được kiểm tra bằng phiên bản Python, Tcl/Tk, pyodbc và thông tin máy chủ SQL trong lần chạy thực tế. Các lệnh khởi động và kiểm thử được thực hiện bằng PowerShell trên Windows.')
p('main.py tạo CoffeeApp, thiết lập giao diện rồi mở LoginWindow. Khi xác thực thành công, CoffeeService được truyền cho MainWindow. Màn hình chính tạo các trang theo vai trò. database_setup.py dùng để tạo cấu trúc và dữ liệu ban đầu, chạy lại không xóa dữ liệu đang có hoặc đặt lại mật khẩu.')
h('Vận dụng tám nội dung học phần')
course1=[r[:] for r in old['tables'][9]]
course1[-1][-1]='price, phone, report_period và các phương thức CoffeeService chia nhiệm vụ xử lý rõ ràng.'
table('Đối chiếu nội dung học phần từ 1 đến 4',course1,[6.4,2.3,6.8])
p('Ví dụ, hàm phone() trong models/service.py xử lý chuỗi số điện thoại. Hàm loại khoảng trắng, cho phép để trống và từ chối chuỗi sai định dạng. Đoạn trích dưới đây là mã hiện đang chạy.')
code(extract('models/service.py','phone'))
course2=[r[:] for r in old['tables'][10]]
course2[3][2]='open mở tập tin với encoding="utf-8-sig", csv.writer ghi báo cáo. Nhật ký lỗi chỉ được tạo khi có lỗi và có giới hạn kích thước.'
table('Đối chiếu nội dung học phần từ 5 đến 8',course2,[6.4,2.3,6.8])
p('Giỏ hàng và quantities minh họa việc dùng từ điển kết hợp vòng lặp. Danh sách bản ghi được lọc, sắp xếp rồi đưa vào Treeview. Kết quả này thể hiện xử lý dữ liệu bằng các cấu trúc cơ bản thay vì phụ thuộc vào thư viện phân tích dữ liệu.')
h('Xử lý giao dịch và bảo vệ dữ liệu')
p('Mỗi thao tác ghi mở một kết nối và thực hiện các bước trong khối try. Nếu thành công thì commit(), nếu có ngoại lệ thì rollback(), cuối cùng close() để giải phóng kết nối. Trong pyodbc, việc xác nhận hoặc hoàn tác giao dịch cần thực hiện trên kết nối tương ứng (mkleehammer và cộng sự, không ghi năm a). Hàm finish_order() dưới đây thể hiện trực tiếp cách xử lý này.')
finish=extract('models/service.py','finish_order')
code(finish)
p('Tham số truy vấn được truyền bằng dấu hỏi để dữ liệu nhập không trở thành câu lệnh SQL. Tên bảng và cột chỉ được lấy từ danh sách ENTITIES cố định. Các thao tác sửa đơn sử dụng PhienBan để ngăn ghi đè dữ liệu cũ. Khóa giao dịch và chỉ mục duy nhất ngăn chiếm cùng bàn hoặc thanh toán một đơn hai lần.')
p('Mật khẩu được dẫn xuất bằng PBKDF2 (Password-Based Key Derivation Function 2), sử dụng SHA-256, dữ liệu ngẫu nhiên riêng và 260.000 vòng lặp trong security.py. Khi kiểm tra, chương trình dùng hmac.compare_digest. Báo cáo mô tả đúng cơ chế đang dùng, không xem đây là khẳng định hệ thống đã được kiểm định an toàn toàn diện (Python Software Foundation, không ghi năm c).')
h('Giao diện và quy trình thao tác')
h('Đăng nhập',3)
p('Màn hình tách phần nhận diện và biểu mẫu. Tên đăng nhập, mật khẩu và nút xác nhận được sắp xếp theo chiều dọc. Lỗi được hiển thị ngay trên biểu mẫu, mật khẩu được xóa sau khi xác thực thành công.')
figure('sodo/dang_nhap.png','Giao diện đăng nhập',maxh=11)
h('Tổng quan và thực đơn',3)
p('Tổng quan hiển thị doanh thu hôm nay, số hóa đơn đã thanh toán, số đơn đang phục vụ và số bàn trống. Biểu đồ lấy dữ liệu thật từ các đơn thử trong cơ sở dữ liệu. Ảnh chụp được thực hiện trước khi thanh toán đơn 80.000 đồng, vì vậy doanh thu hôm nay ở ảnh là 85.000 đồng và có một đơn đang phục vụ.')
figure('sodo/tong_quan.png','Giao diện tổng quan của quản lý',maxh=11.5)
p('Trang Thực đơn cho phép tìm kiếm, thêm, sửa, xóa và ngừng bán sản phẩm theo quyền. Nút sửa và xóa bị vô hiệu hóa khi chưa chọn một dòng. Những sản phẩm đã được chi tiết đơn tham chiếu không được xóa để giữ lịch sử.')
figure('sodo/thuc_don.png','Giao diện quản lý thực đơn',maxh=11.5)
h('Gọi món và thanh toán',3)
p('Màn hình gọi món hiển thị bàn và khách ở trên, thực đơn cùng giỏ ở giữa. Các nút tăng, giảm, áp dụng số lượng và bỏ món nằm cạnh giỏ. Với dữ liệu thử gồm hai cà phê đen và một cà phê sữa, tổng là 80.000 đồng. Dòng “3 món” trên giao diện là tổng số lượng ba ly, giỏ gồm hai loại sản phẩm.')
figure('sodo/goi_mon.png','Giao diện gọi món với tổng đơn 80.000 đồng',maxh=12)
p('Màn hình thu tiền đặt tổng phải thu, ô tiền khách đưa và tiền thừa theo thứ tự kiểm tra tại quầy. Khi số tiền chưa hợp lệ hoặc chưa đủ, nút hoàn tất bị vô hiệu hóa. Ảnh ghi nhận trường hợp nhận 100.000 đồng và trả lại 20.000 đồng.')
figure('sodo/thanh_toan.png','Giao diện xác nhận thu tiền',maxh=12.5,width=10)
h('Báo cáo và giao diện nhân viên',3)
p('Quản lý chọn ngày hoặc tháng rồi xem doanh thu, số hóa đơn, trung bình mỗi đơn và món bán chạy. Kỳ 14/10/2026 đến 20/10/2026 trong lần chạy thử có 8 đơn đã thanh toán, doanh thu 1.200.000 đồng và trung bình 150.000 đồng mỗi đơn.')
figure('sodo/bao_cao_doanh_thu.png','Giao diện báo cáo doanh thu theo kỳ',maxh=11.5)
p('Nhân viên chỉ có ba trang phục vụ. Sơ đồ bàn hiển thị trạng thái bằng màu và chữ. Nút Gọi món mở đơn mới trên bàn trống, còn bàn đang phục vụ cho phép tiếp tục đơn hiện có. Ảnh chụp sau khi thanh toán cho thấy cả tám bàn đều trống.')
figure('sodo/giao_dien_nhan_vien.png','Giao diện sơ đồ bàn của nhân viên',maxh=11.5)

ch(4,'Đánh giá kiểm thử')
h('Mục tiêu và phương pháp')
p('Kiểm thử nhằm đối chiếu các yêu cầu nhập liệu, tính toán, phân quyền và lưu trữ với hành vi thực tế. Bộ tests/test_coffee.py sử dụng unittest, SQL Server thật và các cửa sổ Tkinter. Cơ sở dữ liệu thử được tạo riêng và xóa khi kết thúc, dữ liệu của quán không được dùng làm nơi ghi kết quả thử.')
p('Bộ kiểm thử tích hợp gồm 19 phương thức, có trường hợp dùng nhiều bộ dữ liệu. Một số phép thử giao diện thay hộp thông báo bằng unittest.mock để không chờ thao tác người dùng. Phần truy vấn, ràng buộc và giao dịch chạy trên SQL Server thật. Riêng lần trình diễn bổ sung dùng các cửa sổ thật và hộp thoại lưu CSV thật, không thay thế hộp thoại xuất tập tin.')
h('Kết quả bộ kiểm thử tích hợp')
p('Ngày 20/10/2026, bộ kiểm thử trước khi dọn mã đạt 19 kiểm thử trong 18,003 giây. Sau khi thay cách quản lý giao dịch và xử lý ngoại lệ bằng khối lệnh trực tiếp, lần chạy lại đạt 19 kiểm thử trong 16,156 giây. Hai thời gian chỉ ghi nhận lần chạy tương ứng, không dùng để kết luận chương trình tăng tốc.')
code('python -B -m unittest discover -s tests -v\nRan 19 tests in 16.156s\nOK')
for idx,title in [(11,'Kết quả kiểm thử dữ liệu và nghiệp vụ'),(12,'Kết quả kiểm thử giao dịch, báo cáo và giao diện')]:
    rows=[r[:] for r in old['tables'][idx]]
    rows[0]=['Thứ tự','Trường hợp và kết quả mong đợi','Kết quả thực tế']
    rows=[r[:] for r in rows]
    rows=[r if i==0 else [r[0],r[1].replace('CRUD','Thêm, đọc, sửa, xóa').replace('salt','dữ liệu ngẫu nhiên').replace('disabled','vô hiệu hóa'), 'Đạt'] for i,r in enumerate(rows)]
    table(title,rows,[1.7,9.8,4])
h('Kịch bản chạy thật và xuất tập tin')
p('Lần chạy bổ sung tạo dữ liệu thử riêng gồm năm món khởi đầu, tám bàn, một khách và một nhân viên có tài khoản. Các đơn có thời điểm thanh toán trong bảy ngày được chuẩn bị để kiểm tra báo cáo. Số liệu ở đây là dữ liệu phục vụ kiểm thử, không phải doanh thu của một quán đang kinh doanh.')
demo=[['Chức năng và dữ liệu đầu vào','Kết quả mong đợi','Kết quả thực tế'],
 ['Đăng nhập sai mật khẩu.','Từ chối và giữ biểu mẫu.','Báo lỗi trên biểu mẫu đăng nhập.'],
 ['Giá sản phẩm nhập chữ.','Không lưu sản phẩm.','Biểu mẫu giữ lại và hiển thị lỗi.'],
 ['Thêm món hợp lệ giá 27.000 đồng, sau đó xóa.','Ghi và xóa đúng dữ liệu.','Đối chiếu SQL xác nhận cả hai thao tác.'],
 ['Nhập chữ vào số lượng.','Từ chối dữ liệu sai.','Giỏ không nhận số lượng sai.'],
 ['Hai cà phê đen giá 25.000 đồng, một cà phê sữa giá 30.000 đồng.','Tổng 80.000 đồng.','Giỏ và đơn SQL cùng có tổng 80.000 đồng.'],
 ['Tiền nhận 50.000 đồng, sau đó 100.000 đồng.','Chặn khi thiếu, thừa 20.000 đồng khi đủ.','Nút bị khóa khi thiếu. Khi hoàn tất, khách tăng 8 điểm.'],
 ['Hủy rồi xóa một đơn chưa thanh toán.','Bàn trống, không có doanh thu.','Trạng thái và kết quả tổng hợp đúng.'],
 ['Kỳ 14/10/2026 đến 20/10/2026, xuất CSV.','Tập tin khớp số liệu hiển thị.','8 đơn, 1.200.000 đồng, đúng mã hóa tiếng Việt.'],
 ['Đăng nhập bằng tài khoản nhân viên.','Chỉ có ba trang phục vụ.','Có Đơn hàng, Sơ đồ bàn và Thực đơn.'],
 ['Đóng cửa sổ, đăng nhập và mở lại bằng kết nối mới.','Đơn và điểm đã ghi vẫn còn.','Đơn đã thanh toán 80.000 đồng, khách có 8 điểm.']]
table('Kết quả chạy thật trên giao diện và SQL Server',demo,[6,4.5,5])
p('Tập tin do hộp thoại lưu tạo ra là doanh_thu_2026-10-14.csv, nằm trong thư mục D:\\NCTU\\Python\\ThucHanh\\SoLieuCSVProject do người dùng chọn. Khi đọc lại, tập tin có dấu nhận biết UTF-8, đúng tiêu đề và hai dòng sản phẩm. Cà phê đen có 30 ly, doanh thu 750.000 đồng. Cà phê sữa có 15 ly, doanh thu 450.000 đồng. Tổng hai dòng bằng tổng doanh thu 1.200.000 đồng và số đơn ở đầu tập tin bằng 8.')
h('Đánh giá và giới hạn kiểm tra')
p('Các chức năng đã thử đáp ứng quy trình phục vụ, quản trị và báo cáo trong phạm vi đề tài. Dữ liệu hợp lệ được ghi, dữ liệu sai bị từ chối, các thao tác liên quan được hoàn tác khi có lỗi. Giao diện phân biệt vai trò, làm rõ tổng tiền và trạng thái bàn. Mã sau khi dọn vẫn vượt qua bộ kiểm thử tích hợp đã có.')
p('Việc lưu trữ đã được kiểm tra sau khi đóng cửa sổ và tạo kết nối mới. Chưa thực hiện tắt máy rồi khởi động lại hoặc mô phỏng mất điện. Vì vậy, báo cáo chưa kết luận đạt tình huống tắt máy theo nghĩa đầy đủ. Chưa có thử nghiệm tải lớn, nhiều máy vận hành dài ngày hoặc đánh giá bảo mật độc lập. Những giới hạn này cần được phân biệt với các kết quả đã quan sát.')

front('Phần kết luận')
pp=doc.add_heading('Kết quả đạt được',2)
p('Đề tài đã xây dựng ứng dụng quản lý quán cà phê bằng Python với giao diện Tkinter/ttk, kết nối SQL Server và xuất báo cáo CSV. Chương trình có quản lý danh mục, tài khoản, bàn, khách hàng, đơn hàng, thanh toán và thống kê theo kỳ. Các thao tác quan trọng kiểm tra dữ liệu, quyền và trạng thái trước khi cập nhật.')
p('Kết quả chạy thử xác nhận tính tổng 80.000 đồng, tiền thừa 20.000 đồng, cộng 8 điểm và tập tin báo cáo có 8 đơn với doanh thu 1.200.000 đồng. Bộ kiểm thử tích hợp 19 trường hợp đạt sau khi dọn mã. Tám nội dung học phần được đối chiếu với các tệp và hàm hiện có, giúp giải thích quá trình vận dụng kiến thức vào một ứng dụng hoàn chỉnh trong phạm vi môn học.')
pp=doc.add_heading('Hướng phát triển',2)
p('Trong phạm vi tiếp theo, có thể bổ sung in hóa đơn, quản lý nguyên liệu, sao lưu và phục hồi dữ liệu. Trước khi mở rộng cần hoàn thành thử nghiệm sau khi khởi động lại máy, thử tải và đánh giá quyền truy cập cơ sở dữ liệu. Các chức năng này là định hướng, chưa được xem là kết quả đã cài đặt.')

front('Tài liệu tham khảo')
refs=[
 ('abdelbend','không ghi năm a','db.py trong Restaurant Management System','https://github.com/abdelbend/Restaurant-Management-System/blob/main/db.py'),
 ('abdelbend','không ghi năm b','Restaurant Management System','https://github.com/abdelbend/Restaurant-Management-System'),
 ('Microsoft','không ghi năm','rowversion (Transact-SQL)','https://learn.microsoft.com/en-us/sql/t-sql/data-types/rowversion-transact-sql'),
 ('mkleehammer và cộng sự','không ghi năm a','Database Transaction Management','https://github.com/mkleehammer/pyodbc/wiki/Database-Transaction-Management'),
 ('mkleehammer và cộng sự','không ghi năm b','pyodbc Python ODBC bridge','https://github.com/mkleehammer/pyodbc'),
 ('Python Software Foundation','không ghi năm a','csv CSV File Reading and Writing','https://docs.python.org/3/library/csv.html'),
 ('Python Software Foundation','không ghi năm b','decimal Decimal fixed-point and floating-point arithmetic','https://docs.python.org/3/library/decimal.html'),
 ('Python Software Foundation','không ghi năm c','hashlib Secure hashes and message digests','https://docs.python.org/3/library/hashlib.html'),
 ('Python Software Foundation','không ghi năm d','tkinter Python interface to Tcl/Tk','https://docs.python.org/3/library/tkinter.html'),
 ('Python Software Foundation','không ghi năm e','unittest Unit testing framework','https://docs.python.org/3/library/unittest.html')]
p('Tài liệu tiếng Anh')
for author,year,title,url in refs:
    pp=p(author+' ('+year+'). ','References')
    pp.add_run(title).italic=True
    pp.add_run('. Truy cập ngày 04/10/2026. ')
    hyp=OxmlElement('w:hyperlink')
    hyp.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r')
    rp=OxmlElement('w:rPr')
    ff=OxmlElement('w:rFonts')
    ff.set(qn('w:ascii'),'Times New Roman')
    ff.set(qn('w:hAnsi'),'Times New Roman')
    rp.append(ff)
    sz=OxmlElement('w:sz')
    sz.set(qn('w:val'),'26')
    rp.append(sz)
    rr.append(rp)
    tt=OxmlElement('w:t')
    tt.text=url
    rr.append(tt)
    hyp.append(rr)
    pp._p.append(hyp)

front('Phụ lục\nHướng dẫn cài đặt, sử dụng chương trình/giải pháp')
p('Chuẩn bị môi trường','Appendix Heading')
p('Cài Python cho Windows từ trang chính thức python.org nếu máy chưa có. Trong trình cài đặt, giữ thành phần Tcl/Tk và pip, đồng thời chọn thêm Python vào PATH để gọi bằng lệnh python. Sau khi cài, mở PowerShell và kiểm tra phiên bản. Lần chạy của báo cáo sử dụng Python 3.14.3.')
code('python --version\npython -m tkinter')
p('Chuẩn bị SQL Server hoặc SQL Server Express đang hoạt động và trình điều khiển ODBC phù hợp. Tài khoản Windows chạy chương trình cần quyền tạo cơ sở dữ liệu khi thiết lập lần đầu và quyền đọc, ghi dữ liệu khi sử dụng. Cấu hình mặc định của dự án là THINKPADX1\\SQLSEVER và QuanLyQuanCoffee. Nếu dùng máy khác, sửa biến môi trường COFFEE_SQL_SERVER, COFFEE_SQL_DATABASE hoặc cấu hình trong db_connection.py theo máy thực tế.')
p('Cài phụ thuộc và khởi tạo dữ liệu','Appendix Heading')
p('Tại thư mục dự án, thực hiện lần lượt các lệnh sau. requirements.txt chỉ yêu cầu pyodbc. database_setup.py tạo bảy bảng, năm sản phẩm, tám bàn và tài khoản quản lý khởi đầu. Chạy lại không xóa dữ liệu hoặc đặt lại mật khẩu đã đổi.')
code('python -m pip install -r requirements.txt\npython database_setup.py\npython main.py')
p('Đăng nhập lần đầu bằng admin với mật khẩu Admin@123, sau đó đổi mật khẩu cá nhân. Quản lý tạo hồ sơ nhân viên trước khi cấp tài khoản. Không dùng tài khoản khởi đầu làm thông tin công khai khi đưa hệ thống vào sử dụng.')
p('Quy trình sử dụng','Appendix Heading')
for text in [
 '1. Quản lý kiểm tra thực đơn, giá, bàn, khách và tài khoản trước ca phục vụ.',
 '2. Nhân viên mở Sơ đồ bàn, chọn bàn trống để gọi món hoặc tiếp tục đơn đang phục vụ.',
 '3. Chọn món, thay đổi số lượng và chọn khách nếu cần tích điểm. Lưu đơn để tiếp tục phục vụ. Khi đổi bàn phải chọn bàn chưa có đơn đang xử lý.',
 '4. Chọn Lưu & thu tiền hoặc Thu tiền, nhập số tiền đã nhận, kiểm tra tiền thừa rồi xác nhận hoàn tất. Quay lại màn hình gọi món không hủy đơn đã lưu.',
 '5. Quản lý mở Báo cáo doanh thu, chọn kỳ, xem số liệu rồi dùng Xuất CSV và chọn nơi lưu. Mở tập tin để đối chiếu tổng doanh thu và số đơn.',
 '6. Dùng F5 để làm mới trang hiện tại. Đăng xuất khi kết thúc phiên làm việc.']:
    p(text)
p('Chạy kiểm thử và chuẩn bị trình bày','Appendix Heading')
code('python -B -m unittest discover -s tests -v')
p('Bộ kiểm thử cần SQL Server và cửa sổ Tkinter. Cơ sở dữ liệu thử được tạo riêng rồi xóa sau khi hoàn tất. Kịch bản trình bày và cách giải thích giỏ hàng, giá lúc bán, giao dịch được lưu trong README.md.')
in_appendix=False
for pp in doc.paragraphs:
    if pp.text.startswith('PHỤ LỤC'):in_appendix=True
    if in_appendix:pp.paragraph_format.line_spacing=1.0
update=OxmlElement('w:updateFields')
update.set(qn('w:val'),'true')
doc.settings.element.append(update)
doc.core_properties.author='Lê Minh Phú'
doc.core_properties.title='Báo cáo học phần Lập trình Python - Quản lý quán cà phê'
doc.core_properties.subject='Ứng dụng Tkinter và SQL Server'
doc.core_properties.comments=''
doc.save(OUT)
(W/'inventory.json').write_text(json.dumps({'figures':figures,'tables':tables},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved',OUT,len(figures),'figures',len(tables),'tables')
