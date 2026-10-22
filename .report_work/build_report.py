from pathlib import Path
import json, math, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

W = Path(__file__).parent
ROOT = W.parent
ENV = json.loads((W/'evidence.json').read_text(encoding='utf-8'))
FONT = 'C:/Windows/Fonts/times.ttf'
BOLD = 'C:/Windows/Fonts/timesbd.ttf'

class Diagram:
    def __init__(self, name, size=(1800,1100)):
        self.name=name
        self.im=Image.new('RGB',size,'white')
        self.d=ImageDraw.Draw(self.im)
    def text(self,xy,text,size=32,bold=False):
        self.d.multiline_text(xy,text,font=ImageFont.truetype(BOLD if bold else FONT,size),fill='black',anchor='mm',align='center',spacing=8)
    def box(self,b,text,ellipse=False):
        (self.d.ellipse if ellipse else self.d.rectangle)(b,outline='black',width=3,fill='white')
        self.text(((b[0]+b[2])/2,(b[1]+b[3])/2),text)
    def arrow(self,pts,label=None,pos=None):
        self.d.line(pts,fill='black',width=3)
        x,y=pts[-1]
        x0,y0=pts[-2]
        a=math.atan2(y-y0,x-x0)
        self.d.polygon([(x,y),(x-20*math.cos(a-.45),y-20*math.sin(a-.45)),(x-20*math.cos(a+.45),y-20*math.sin(a+.45))],fill='black')
        if label:
            self.text(pos or ((pts[0][0]+x)/2,(pts[0][1]+y)/2-26),label,29)
    def save(self):
        self.im.save(W/f'{self.name}.png')

# DFD ngữ cảnh, toàn hệ thống là một tiến trình.
d=Diagram('dfd',(1800,1120))
d.box((40,330,400,580),'Nhân viên\nQuản lý')
d.box((650,330,1130,620),'0\nHệ thống quản lý\nquán cà phê',True)
d.box((1400,330,1760,580),'Người nhận\nbáo cáo')
d.arrow([(400,370),(650,370)],'Thông tin đăng nhập,\ndanh mục, đơn hàng',(525,240))
d.arrow([(650,540),(400,540)],'Kết quả xử lý,\ntrạng thái bàn',(525,680))
d.arrow([(1130,410),(1400,410)],'Báo cáo CSV',(1265,330))
d.box((630,860,1150,1020),'D1 · SQL Server\nDanh mục, tài khoản, đơn hàng')
d.d.line([(1150,863),(1150,1017)],fill='white',width=6)
d.arrow([(790,620),(790,860)],'Dữ liệu cập nhật',(610,770))
d.arrow([(1010,860),(1010,620)],'Dữ liệu tra cứu',(1210,770))
d.text((900,70),'DFD CẤP 0',40,True)
d.save()

d=Diagram('functions',(1800,1050))
d.box((600,40,1200,170),'QUẢN LÝ QUÁN CÀ PHÊ')
groups=[('Phiên làm việc',['Đăng nhập','Đổi mật khẩu','Đăng xuất']),('Danh mục',['Món, bàn','Khách, nhân viên','Tài khoản / vai trò']),('Phục vụ',['Gọi món, sửa đơn','Thu tiền / huỷ đơn','Tính điểm, trả bàn']),('Báo cáo',['Tổng quan','Lọc ngày / tháng','Biểu đồ, xuất CSV'])]
d.arrow([(900,170),(900,260)])
d.d.line([(225,260),(1575,260)],fill='black',width=3)
for i,(title,items) in enumerate(groups):
    x=225+i*450
    d.arrow([(x,260),(x,340)])
    d.box((x-190,340,x+190,450),title)
    for j,t in enumerate(items):
        yy=560+j*160
        d.d.line([(x,450),(x-210,450),(x-210,yy+52)],fill='black',width=3)
        d.arrow([(x-210,yy+52),(x-190,yy+52)])
        d.box((x-190,yy,x+190,yy+105),t)
d.save()

d=Diagram('architecture',(1800,800))
for b,t in [((70,200,500,430),'Giao diện Tkinter\nui/ và main.py'),((685,200,1115,430),'Nghiệp vụ\nmodels/service.py'),((1300,200,1730,430),'Dữ liệu SQL Server\ndb_connection.py')]:d.box(b,t)
d.arrow([(500,260),(685,260)],'Yêu cầu',(590,180))
d.arrow([(685,380),(500,380)],'Kết quả',(590,460))
d.arrow([(1115,260),(1300,260)],'SQL + tham số',(1208,180))
d.arrow([(1300,380),(1115,380)],'Bản ghi',(1208,460))
d.box((70,590,500,735),'Tệp CSV\nui/statistics_tab.py')
d.arrow([(285,430),(285,590)],'Ghi báo cáo',(475,510))
d.text((900,65),'TỔ CHỨC 3 PHẦN TRONG CÙNG ỨNG DỤNG',38,True)
d.save()

d=Diagram('er',(1800,1050))
boxes={'nv':((50,100,430,280),'NhanVien\nPK MaNV'), 'tk':((50,560,430,790),'TaiKhoan\nPK MaTK\nFK MaNV (UNIQUE)'), 'kh':((710,70,1090,270),'KhachHang\nPK MaKH'), 'ban':((1370,70,1750,270),'Ban\nPK MaBan'), 'dh':((710,450,1090,690),'DonHang\nPK MaDH\nFK MaNV, MaKH, MaBan'), 'ct':((710,850,1090,1020),'ChiTietDonHang\nPK MaCTDH\nFK MaDH, MaSP'), 'sp':((1370,830,1750,1020),'SanPham\nPK MaSP')}
for b,t in boxes.values():d.box(b,t)
for a,b,l,p in [([(240,280),(240,560)],None,'1 → 0..1',(390,425)), ([(430,190),(560,190),(560,530),(710,530)],None,'1 → N',(555,360)), ([(900,270),(900,450)],None,'0..1 ← N',(1040,360)), ([(1560,270),(1560,570),(1090,570)],None,'1 → N',(1350,500)), ([(900,690),(900,850)],None,'1 → N',(1040,770)), ([(1370,930),(1090,930)],None,'1 → N',(1230,860))]:d.arrow(a,l,p)
d.save()

def flow(name, steps, decision_index, no_text, back=False):
    d=Diagram(name,(1600,1500))
    x=770
    for i,t in enumerate(steps):
        y=70+i*205
        if i==decision_index:
            d.d.polygon([(x,y),(x+300,y+80),(x,y+160),(x-300,y+80)],fill='white',outline='black',width=3)
            d.text((x,y+80),t,31)
            d.arrow([(x+300,y+80),(1380,y+80),(1380,y+330)],'Không',(1235,y+40))
            d.box((1160,y+330,1580,y+470),no_text)
            if back:
                d.arrow([(1380,y+470),(1380,1485),(180,1485),(180,330),(470,330)])
            else:
                d.arrow([(1380,y+470),(1380,1380),(1070,1380)])
        else:
            d.box((470,y,1070,y+160),t,ellipse=i in (0,len(steps)-1))
        if i<len(steps)-1:d.arrow([(x,y+160),(x,y+205)],'Có' if i==decision_index else None,(x-70,y+180))
    d.save()
flow('flow_order',['Bắt đầu lưu đơn','Gộp món trùng, kiểm tra\nsố lượng và dữ liệu','Quyền, bàn, món và\nphiên bản hợp lệ?','Lặp từng món, lấy giá\nlịch sử và tính tổng','Ghi đơn, chi tiết, cập nhật bàn\ntrong một giao dịch','Commit và trả mã đơn','Kết thúc'],2,'Báo lỗi\nRollback nếu đã mở\ngiao dịch')
flow('flow_payment',['Bắt đầu thu tiền','Nhập tiền khách đưa\nTính tiền thừa','Tiền khách đưa\n≥ tổng đơn?','Kiểm tra đơn đang xử lý\nvà phiên bản hiện tại','Ghi thanh toán, cộng điểm\ncập nhật trạng thái bàn','Commit\nLàm mới giao diện','Kết thúc'],2,'Hiện tiền còn thiếu\nKhoá nút xác nhận',True)

doc=Document()
sec=doc.sections[0]
def geometry(sec):
    sec.page_width=Cm(21)
    sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2)
    sec.bottom_margin=Cm(2)
    sec.left_margin=Cm(3.5)
    sec.right_margin=Cm(2)
    sec.footer_distance=Cm(1)
geometry(sec)
styles=doc.styles
for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption','TOC Heading']:
    s=styles[name]
    s.font.name='Times New Roman'
    s.font.size=Pt(13)
    s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if s._element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
    rf=s._element.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
    for a in list(rf.attrib):
        if 'theme' in a.lower():del rf.attrib[a]
    s.paragraph_format.line_spacing=1.2
    s.paragraph_format.space_after=Pt(6)
normal=styles['Normal']
normal.paragraph_format.first_line_indent=Cm(1)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for level in (1,2,3):
    s=styles[f'Heading {level}']
    s.font.bold=True
    s.paragraph_format.first_line_indent=Cm(0)
    s.paragraph_format.space_before=Pt(10)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
styles['Heading 1'].font.size=Pt(15)
styles['Caption'].font.italic=True
styles['Caption'].paragraph_format.first_line_indent=Cm(0)
styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ('TOC 1','TOC 2','TOC 3'):
    if name not in styles:
        styles.add_style(name,1)
    styles[name].font.name='Times New Roman'
    styles[name].font.size=Pt(13)
    styles[name].paragraph_format.line_spacing=1.2
    styles[name].paragraph_format.space_after=Pt(3)
    styles[name].paragraph_format.first_line_indent=Cm(0)
doc.core_properties.author='Lê Minh Phú'
doc.core_properties.title='Quản lý quán cà phê bằng Python'
doc.core_properties.subject='Báo cáo học phần Lập trình Python - Thực hành'
doc.core_properties.keywords='Python, Tkinter, SQL Server, quản lý quán cà phê'
doc.core_properties.last_modified_by='Lê Minh Phú'
doc.core_properties.comments=''

def field(p,instruction):
    r=p.add_run()
    for tag in ('begin','separate','end'):
        el=OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'),tag)
        if tag=='begin':
            r._r.append(el)
            text=OxmlElement('w:instrText')
            text.set(qn('xml:space'),'preserve')
            text.text=' '+instruction+' '
            r._r.append(text)
        else:r._r.append(el)
    return r
def p(text):
    assert ';' not in text
    return doc.add_paragraph(text)
def h(text,level=2):return doc.add_heading(text,level)
def page(title,level=2):
    pp=h(title,level)
    pp.paragraph_format.page_break_before=True
def caption(label,text):
    pp=doc.add_paragraph(style='Caption')
    pp.add_run(label+' ')
    field(pp,'SEQ '+('Figure' if label=='Hình' else 'Table')+' \\* ARABIC')
    pp.add_run('. '+text)
def picture(name,desc,width=15.5):
    pp=doc.add_paragraph()
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.keep_with_next=True
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=pp.add_run()
    shape=run.add_picture(str(W/f'{name}.png'),width=Cm(width))
    shape._inline.docPr.set('descr',desc)
    caption('Hình',desc)
def table(title,headers,rows,widths=None):
    caption('Bảng',title)
    t=doc.add_table(rows=1, cols=len(headers))
    t.autofit=False
    widths=widths or [15.5/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit')
        trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for c,txt,w in zip(cells,row,widths):
            c.width=Cm(w)
            c.text=str(txt)
            mar=OxmlElement('w:tcMar')
            for edge,n in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+edge)
                e.set(qn('w:w'),n)
                e.set(qn('w:type'),'dxa')
                mar.append(e)
            c._tc.get_or_add_tcPr().append(mar)
            for pp in c.paragraphs:
                pp.paragraph_format.first_line_indent=Cm(0)
                pp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                pp.paragraph_format.space_after=Pt(0)
                pp.paragraph_format.line_spacing=1.2
                for r in pp.runs:r.bold=i==0
    return t
def code(text):
    pp=doc.add_paragraph()
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.left_indent=Cm(.3)
    pp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pp.paragraph_format.keep_together=True
    r=pp.add_run(text)
    r.font.name='Times New Roman'
    r.font.size=Pt(13)

# Bìa theo mẫu của học phần.
sec.different_first_page_header_footer=True
b=OxmlElement('w:pgBorders')
b.set(qn('w:offsetFrom'),'page')
for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge)
    for k,v in [('val','single'),('sz','12'),('space','24'),('color','000000')]:e.set(qn('w:'+k),v)
    b.append(e)
sec._sectPr.append(b)
def cover(text,size=13,bold=False,color=None,after=12):
    pp=doc.add_paragraph()
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.space_after=Pt(after)
    rr=pp.add_run(text)
    rr.bold=bold
    rr.font.size=Pt(size)
    if color:rr.font.color.rgb=RGBColor.from_string(color)
    return pp
cover('TRƯỜNG ĐẠI HỌC NAM CẦN THƠ',14,True,after=3)
cover('TRƯỜNG CÔNG NGHỆ SỐ VÀ TRÍ TUỆ NHÂN TẠO',13,True,after=16)
pp=cover('',after=22)
pp.add_run().add_picture(str(W/'school_logo.png'),width=Cm(3.6))
cover('BÁO CÁO HỌC PHẦN',20,True,'173F8A',10)
cover('LẬP TRÌNH PYTHON - THỰC HÀNH',18,True,'C00000',30)
cover('Đề tài',14,False,after=8)
cover('XÂY DỰNG ỨNG DỤNG\nQUẢN LÝ QUÁN CÀ PHÊ BẰNG PYTHON',18,True,after=10)
cover('Coffee Shop Management Application',14,False,after=35)
cover('Sinh viên thực hiện',14,True,after=8)
cover('Lê Minh Phú - MSSV: 236717\nLớp: DH23TIN02',14,False,after=18)
cover('Giảng viên hướng dẫn\nBùi Thị Diễm Trinh',14,False,after=45)
cover('Cần Thơ, tháng 10 năm 2026',13,False,after=0)

front=doc.add_section(WD_SECTION_START.NEW_PAGE)
geometry(front)
front.different_first_page_header_footer=False
front._sectPr.remove(front._sectPr.find(qn('w:pgBorders')))
front.footer.is_linked_to_previous=False
pg=OxmlElement('w:pgNumType')
pg.set(qn('w:fmt'),'lowerRoman')
pg.set(qn('w:start'),'1')
front._sectPr.append(pg)
fp=front.footer.paragraphs[0]
fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent=Cm(0)
field(fp,'PAGE')
cover('MỤC LỤC',15,True,after=12)
field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')
doc.add_page_break()
cover('DANH MỤC HÌNH ẢNH VÀ SƠ ĐỒ',15,True,after=12)
field(doc.add_paragraph(),'TOC \\h \\z \\c "Figure"')
doc.add_page_break()
cover('DANH MỤC BẢNG',15,True,after=12)
field(doc.add_paragraph(),'TOC \\h \\z \\c "Table"')

body=doc.add_section(WD_SECTION_START.NEW_PAGE)
geometry(body)
body.footer.is_linked_to_previous=False
pg=body._sectPr.find(qn('w:pgNumType'))
pg.set(qn('w:fmt'),'decimal')
pg.set(qn('w:start'),'1')
fp=body.footer.paragraphs[0]
fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent=Cm(0)
field(fp,'PAGE')
h('PHẦN GIỚI THIỆU',1)
h('Đặt vấn đề')
p('Một quán cà phê nhỏ cần theo dõi bàn, món đã gọi, tiền phải thu và doanh thu cuối ngày. Khi ghi chép rời rạc, nhân viên có thể bỏ sót món, dùng nhầm giá hoặc không biết bàn nào đang phục vụ. Người quản lý cũng mất thời gian cộng lại các hoá đơn để kiểm tra hoạt động của quán.')
p('Đề tài xây dựng ứng dụng desktop Coffee bằng Python để thực hiện một quy trình liên tục từ chọn bàn, gọi món đến thanh toán và xem báo cáo. Giao diện sử dụng Tkinter/ttk, dữ liệu được lưu trong SQL Server, báo cáo có thể xuất ra CSV. Quy mô được giới hạn ở một quán để mã nguồn dễ đọc và các thao tác dễ trình bày trong buổi bảo vệ.')
h('Mục tiêu của đề tài')
p('Ứng dụng cần lưu và tra cứu được sản phẩm, bàn, khách hàng, nhân viên và tài khoản. Quy trình bán hàng phải tính đúng tổng tiền, giữ giá đã bán, ghi nhận thanh toán, cộng điểm khách hàng và tự cập nhật trạng thái bàn. Quản lý cần xem doanh thu theo thời gian và xuất số liệu để kiểm tra bằng phần mềm bảng tính.')
p('Về giao diện, thao tác phục vụ được tập trung vào sơ đồ bàn và màn hình gọi món. Các nút quan trọng có màu và nhãn rõ ràng. Về kỹ thuật, phần giao diện, xử lý nghiệp vụ và kết nối dữ liệu được tách thành các tệp nhỏ, tránh đưa toàn bộ chương trình vào một cửa sổ hoặc một hàm dài.')
h('Đối tượng và phạm vi sử dụng')
p('Nhân viên sử dụng ứng dụng để xem thực đơn, chọn bàn, tạo và sửa đơn đang xử lý, thu tiền và huỷ đơn. Quản lý có thêm quyền quản lý danh mục, nhân sự, tài khoản và báo cáo. Đây là ứng dụng desktop trên Windows, sử dụng chuột và bàn phím, không có giao diện web hay ứng dụng di động.')
p('Phạm vi hiện tại chưa bao gồm quản lý nguyên liệu, giảm giá, thanh toán qua ngân hàng, hoá đơn điện tử và nhiều chi nhánh. CSV được dùng để xuất báo cáo. Chương trình chưa có chức năng nhập dữ liệu từ tệp vào danh mục.')

page('Những nghiên cứu và thư viện liên quan')
p('Các tài liệu dưới đây được khảo sát để đối chiếu cách tổ chức ứng dụng Python và lựa chọn công cụ phù hợp với đề tài. Nguồn tham khảo gồm một ứng dụng quản lý nhà hàng và các thư viện có sử dụng trực tiếp trong chương trình.')
table('Ứng dụng và thư viện tham khảo',['Nguồn','Nội dung tham khảo','Áp dụng trong đề tài'],[
['Restaurant Management System [1]','Ứng dụng Python quản lý thực đơn và đơn hàng.','Đối chiếu quy trình chọn món và xem hoá đơn.'],
['Tkinter/ttk [2]','Widget, biến giao diện và vòng lặp sự kiện.','Cửa sổ, bảng, hộp thoại, nút và bộ lọc.'],
['pyodbc [3]','Cầu nối Python và ODBC.','Kết nối SQL Server, truyền tham số truy vấn.'],
['csv [4]','Ghi dữ liệu dạng hàng và cột.','Xuất báo cáo CSV tiếng Việt.'],
['decimal [5]','Số thập phân với độ chính xác kiểm soát được.','Tính tiền và tổng doanh thu.'],
['unittest [6]','Tổ chức và chạy kiểm thử.','Kiểm tra nghiệp vụ, dữ liệu và giao diện.']], [4.1,5.5,5.9])
h('Phân tích ứng dụng tham khảo chính')
p('Restaurant Management System của abdelbend dùng Python và SQLite, có giao diện dòng lệnh và Tkinter. Chương trình tự tạo cơ sở dữ liệu, bổ sung thực đơn mẫu, hỗ trợ quản lý món, đặt món và xem hoá đơn. Các món thuộc cùng đơn được nhóm bằng order_id [1].')
p('Tệp db.py tập trung việc mở kết nối, khởi tạo bảng và thực hiện truy vấn. Dữ liệu chính gồm menu và customer, trong đó customer lưu từng dòng món của đơn. Cách nhóm dòng món giúp dựng hoá đơn, nhưng mô hình này khá gọn và chưa thể hiện riêng bàn, tài khoản hoặc các trạng thái phục vụ của quán [1, 7].')
p('Đề tài Coffee đối chiếu cách tách thao tác dữ liệu và cách cộng tiền theo các dòng món. Giải pháp của đề tài dùng DonHang và ChiTietDonHang để tách thông tin chung khỏi chi tiết, bổ sung bàn và hai vai trò. Một lần lưu hoặc thanh toán được gom trong cùng giao dịch để tránh cập nhật dở dang.')

page('Luồng dữ liệu tổng quát')
p('DFD cấp 0 mô tả tổng quát dữ liệu đi vào và đi ra khỏi ứng dụng. Toàn bộ ứng dụng là tiến trình 0. Nhân viên và quản lý cung cấp dữ liệu đầu vào, nhận kết quả xử lý. SQL Server là kho lưu trữ chính. Báo cáo CSV là dữ liệu đầu ra được giao cho người nhận báo cáo.')
picture('dfd','DFD cấp 0 của ứng dụng Coffee')
p('Luồng cập nhật bao gồm danh mục, thông tin đơn, trạng thái thanh toán và điểm khách hàng. Luồng tra cứu trả về thông tin để hiển thị trên giao diện. Chỉ quản lý được yêu cầu báo cáo doanh thu. CSV không thay thế cơ sở dữ liệu và không phải đầu vào tự động của chương trình.')
p('Ví dụ, khi nhân viên xác nhận đã thu đủ tiền, ứng dụng kiểm tra đơn, ghi trạng thái đã thanh toán, cập nhật điểm khách và trả bàn về trạng thái trống. Khi quản lý chọn một khoảng ngày, chương trình đọc các đơn đã thanh toán trong kỳ rồi hiển thị tổng hợp hoặc ghi tệp CSV.')

page('CHƯƠNG 1. ĐẶC TẢ YÊU CẦU',1)
h('Các yêu cầu chức năng')
table('Các nhóm chức năng chính',['Mã','Chức năng','Kết quả yêu cầu'],[
['CN01','Đăng nhập và phiên làm việc','Xác thực mật khẩu, xác định vai trò, đổi mật khẩu, đăng xuất.'],
['CN02','Quản lý danh mục','Quản lý thêm, tìm, sửa, xoá món, bàn, khách hàng và nhân viên.'],
['CN03','Quản lý tài khoản','Cấp tài khoản, chọn vai trò, khoá/mở và đặt lại mật khẩu.'],
['CN04','Phục vụ bàn','Hiển thị bàn trống hoặc đang phục vụ, mở đúng đơn hiện có.'],
['CN05','Xử lý đơn','Chọn món, đổi số lượng, bỏ món, lưu đơn và chuyển bàn hợp lệ.'],
['CN06','Thanh toán và huỷ','Tính tiền thừa, ghi thu tiền, cộng điểm, trả bàn, giữ lịch sử.'],
['CN07','Thống kê','Tổng doanh thu, số hoá đơn, trung bình và món bán chạy.'],
['CN08','Xuất báo cáo','Xuất CSV cho kỳ đã chọn, giữ tiếng Việt và số liệu đúng.']], [1.5,4.3,9.7])
p('Các màn hình danh mục có ô tìm kiếm và bảng dữ liệu. Quản lý lựa chọn bản ghi trước khi sửa hoặc xoá. Nhân viên được xem thực đơn để phục vụ nhưng không được sửa danh mục. Chương trình kiểm tra quyền ở tầng nghiệp vụ để việc ẩn nút trên giao diện không trở thành cơ chế bảo vệ duy nhất.')

page('Quy tắc nghiệp vụ và kiểm tra dữ liệu')
table('Các quy tắc cần đảm bảo',['Dữ liệu / thao tác','Quy tắc'],[
['Tên và giá món','Tên không được rỗng. Giá nguyên đồng từ 1 đến 1.000.000.000.'],
['Số lượng','Số nguyên từ 1 đến 999 cho mỗi món sau khi gộp các dòng trùng.'],
['Số điện thoại','Có thể để trống. Nếu nhập, dùng 10 số bắt đầu bằng 0 hoặc +84 và 9 số.'],
['Trạng thái bàn','Mỗi bàn có tối đa một đơn đang xử lý. Lưu, chuyển, huỷ và thanh toán cập nhật bàn tương ứng.'],
['Giá lịch sử','Chi tiết đã lưu giữ tên và đơn giá lúc bán. Đổi giá thực đơn không sửa hoá đơn cũ.'],
['Thu tiền','Chỉ xác nhận đơn đang xử lý và nhận đủ tiền. Mỗi 10.000 đồng được 1 điểm, làm tròn xuống.'],
['Lịch sử và liên kết','Không sửa hoặc xoá đơn đã thanh toán. Không xoá bản ghi còn được tham chiếu.'],
['Tài khoản quản lý','Không tự xoá, tự khoá hoặc hạ quyền. Phải còn ít nhất một quản lý hoạt động.']], [4.1,11.4])
h('Yêu cầu về sử dụng và độ tin cậy')
p('Chữ tiếng Việt phải đọc rõ, màu trạng thái phải đi kèm nhãn và các thao tác xoá hoặc huỷ phải có xác nhận. Cửa sổ dùng bố cục co giãn, vùng dài có thanh cuộn. Phím F5 làm mới trang hiện tại. Người sử dụng nhận thông báo cụ thể khi dữ liệu sai hoặc không còn hợp lệ.')
p('Cơ sở dữ liệu phải giữ được dữ liệu giữa các lần chạy. Những thay đổi thuộc cùng nghiệp vụ phải được ghi đầy đủ hoặc hoàn tác. Nếu một đơn đã thay đổi ở cửa sổ khác, bản cũ không được ghi đè. Mật khẩu không được lưu dưới dạng văn bản thuần.')

page('Phân quyền và các tình huống sử dụng')
table('Đối chiếu quyền sử dụng',['Thao tác','Nhân viên','Quản lý'],[
['Xem bàn và thực đơn','Có','Có'],['Tạo, sửa đơn đang xử lý','Có','Có'],['Thanh toán, huỷ đơn đang xử lý','Có','Có'],['Xoá đơn chưa thanh toán','Không','Có, theo trạng thái hợp lệ'],['Sửa danh mục và nhân sự','Không','Có'],['Quản lý tài khoản','Không','Có'],['Xem thống kê, xuất CSV','Không','Có'],['Đổi mật khẩu cá nhân','Có','Có']], [8.5,3.5,3.5])
h('Tình huống phục vụ một bàn')
p('Nhân viên đăng nhập, chọn bàn trống và mở gọi món. Nhân viên chọn khách hàng nếu cần tích điểm, thêm món và kiểm tra giỏ. Lưu đơn chuyển bàn sang đang phục vụ. Trong thời gian phục vụ, nhân viên có thể mở lại đơn để điều chỉnh số lượng hoặc chuyển sang bàn trống.')
p('Khi khách thanh toán, nhân viên mở thu tiền và nhập số tiền đã nhận. Nếu chưa đủ, nút hoàn tất bị vô hiệu hoá. Khi nhận đủ và xác nhận, đơn chuyển sang đã thanh toán, khách có đăng ký được cộng điểm và bàn trở lại trống.')
h('Tình huống cuối ca')
p('Quản lý chọn báo cáo doanh thu, xác định khoảng ngày hoặc tháng và xem số liệu. Chỉ các đơn đã thanh toán trong kỳ được tính. Quản lý đối chiếu tổng tiền, số hoá đơn, các món bán chạy rồi xuất CSV tới vị trí mong muốn. Đơn đang xử lý và đơn huỷ không được tính là doanh thu.')

page('CHƯƠNG 2. THIẾT KẾ GIẢI PHÁP',1)
h('Phân rã chức năng')
p('Hệ thống được chia thành bốn nhóm: phiên làm việc, danh mục, phục vụ và báo cáo. Nhóm phục vụ là luồng chính của nhân viên. Nhóm danh mục và báo cáo giúp quản lý chuẩn bị dữ liệu và kiểm tra kết quả hoạt động.')
picture('functions','Sơ đồ phân rã chức năng')
p('Các chức năng cùng dùng dịch vụ CoffeeService. Cửa sổ gọi món và thu tiền là hộp thoại riêng để người sử dụng tập trung vào đơn hiện tại. Các màn hình danh mục có cấu trúc chung, vì vậy có thể dùng EntityTab và FormDialog để giảm mã lặp.')

page('Kiến trúc và tổ chức mã nguồn')
picture('architecture','Các phần của ứng dụng và hướng trao đổi dữ liệu')
table('Trách nhiệm của các tệp chính',['Thành phần','Trách nhiệm'],[
['main.py, ui/','Khởi động, đăng nhập, điều hướng và thu thập thao tác.'],['models/service.py','Kiểm tra dữ liệu, quyền, xử lý đơn và thống kê.'],['db_connection.py','Mở kết nối, quản lý giao dịch, chuyển bản ghi thành dict.'],['database_setup.py','Tạo cấu trúc dữ liệu và dữ liệu khởi đầu.'],['security.py','Băm và kiểm tra mật khẩu.'],['ui/theme.py, assets/','Màu sắc, kiểu chữ, trạng thái và hình ảnh control.'],['tests/test_coffee.py','Kiểm thử tích hợp bằng SQL Server và Tkinter.']], [6,9.5])
p('Đây là ba phần trách nhiệm trong một ứng dụng desktop, không phải ba dịch vụ độc lập. Không sử dụng ORM, máy chủ web hoặc hệ thống triển khai phân tán. Truy vấn nằm ở tầng nghiệp vụ, còn giao diện chủ yếu hiển thị và gọi hàm xử lý.')

page('Thiết kế cơ sở dữ liệu')
p('Cơ sở dữ liệu gồm bảy bảng. DonHang lưu thông tin chung, ChiTietDonHang lưu từng món. NhanVien được liên kết với TaiKhoan và DonHang. KhachHang có thể không được chọn khi khách mua lẻ. Các mũi tên trong sơ đồ đi từ bản ghi được tham chiếu đến bảng chứa khoá ngoại.')
picture('er','Quan hệ giữa bảy bảng dữ liệu')
p('Một nhân viên có tối đa một tài khoản. Mỗi đơn thuộc một bàn, một nhân viên và có thể thuộc một khách hàng. Một đơn có nhiều chi tiết, mỗi chi tiết tham chiếu một sản phẩm. Cặp MaDH và MaSP là duy nhất để một món không bị tách thành các dòng trùng trong cùng đơn.')
p('Khoá ngoại bảo vệ dữ liệu có liên kết. Chỉ mục duy nhất có điều kiện trên bàn của đơn đang xử lý ngăn hai đơn hoạt động cùng chiếm một bàn. PhienBan dùng kiểu ROWVERSION để phát hiện đơn đã bị thay đổi. ROWVERSION là dấu phiên bản nhị phân, không phải ngày giờ [8].')

page('Các trường dữ liệu quan trọng')
table('Cấu trúc bảng và trường cần giải thích',['Bảng','Trường và ý nghĩa'],[
['NhanVien','MaNV: khoá chính. HoTen, SDT, ChucVu: thông tin nhân sự.'],['TaiKhoan','MaTK, TenDangNhap duy nhất, MatKhauHash, VaiTro, MaNV duy nhất, HoatDong.'],['KhachHang','MaKH, HoTen, SDT, DiemTichLuy không âm.'],['Ban','MaBan, TenBan duy nhất, KhuVuc, TrangThai.'],['SanPham','MaSP, TenSP, DanhMuc, Gia DECIMAL(18,2), MoTa, TrangThai bán.'],['DonHang','MaDH, MaBan, MaKH, MaNV, thời gian tạo/thanh toán, TongTien, DiemDaCong, TrangThai, PhienBan.'],['ChiTietDonHang','MaCTDH, MaDH, MaSP, TenSPLucBan, SoLuong, DonGia và ThanhTien tính từ số lượng × đơn giá.']], [3.7,11.8])
h('Lựa chọn kiểu dữ liệu')
p('Các tên tiếng Việt dùng NVARCHAR. Tiền dùng DECIMAL(18,2) trong SQL Server và Decimal trong Python, dù quy tắc nhập chỉ nhận giá nguyên đồng. Số lượng và điểm dùng số nguyên. Thời gian tạo và thanh toán dùng DATETIME2. Trạng thái tài khoản và bán món dùng BIT.')
p('Tên và giá lúc bán được lưu ngay trong chi tiết đơn để lịch sử vẫn đọc được khi thực đơn đổi. DiemDaCong cho biết điểm đã cộng của đơn. Tài khoản lưu chuỗi băm chứa thông tin cần thiết để kiểm tra mật khẩu, không lưu mật khẩu gốc.')

page('Cấu trúc dữ liệu Python và định dạng CSV')
table('List và dict trong chương trình',['Cấu trúc','Ví dụ và cách dùng'],[
['List các dict','Danh sách sản phẩm và kết quả truy vấn. Mỗi dict có tên cột làm khoá.'],['Dict giỏ hàng','Khoá là MaSP. Giá trị chứa MaSP, TenSP, SoLuong và DonGia.'],['Dict ánh xạ lựa chọn','Nhãn bàn hoặc khách hàng ánh xạ tới mã tương ứng để tránh nhập mã bằng tay.'],['Dict quantities','Gộp số lượng các dòng cùng mã sản phẩm trước khi lưu.'],['List các tuple','Dãy thời gian và doanh thu cung cấp cho biểu đồ.']], [4.4,11.1])
p('Dùng dict theo mã sản phẩm giúp tìm món trong giỏ nhanh và cập nhật đúng số lượng. List giữ các dòng dữ liệu theo thứ tự để đưa lên bảng hoặc ghi tệp. Các cấu trúc này đủ cho quy mô đề tài, không cần thêm thư viện xử lý dữ liệu.')
table('Cấu trúc tệp CSV xuất báo cáo',['Dòng','Nội dung'],[
['1','Từ ngày, giá trị ngày đầu, Đến ngày, giá trị ngày cuối.'],['2','Tổng doanh thu (đồng), giá trị tổng, Đơn đã thanh toán, số đơn.'],['3','Dòng trống phân cách thông tin kỳ và dữ liệu món.'],['4','Mã sản phẩm, Sản phẩm, Số lượng, Doanh thu (đồng).'],['5 trở đi','Một dòng cho mỗi sản phẩm thuộc kết quả thống kê bán chạy.']], [2.4,13.1])
p('CSV dùng UTF-8 có BOM và dấu phẩy do csv.writer quản lý. newline="" tránh dòng trống ngoài ý muốn. Tên món bắt đầu bằng =, +, - hoặc @ được thêm dấu nháy đơn để giảm khả năng bị phần mềm bảng tính hiểu là công thức. Tệp hiện phục vụ xuất và đối chiếu báo cáo, không dùng làm nguồn nhập danh mục [4].')

page('Thuật toán lưu đơn hàng')
picture('flow_order','Lưu đồ kiểm tra và lưu đơn hàng',12.4)
p('Vòng lặp gộp số lượng theo mã sản phẩm trước khi lấy giá và tính thành tiền. Khi sửa, đơn phải còn đang xử lý và đúng phiên bản. Nếu món ngừng bán, không được thêm mới hoặc tăng số lượng. Toàn bộ ghi đơn, chi tiết và trạng thái bàn nằm trong một giao dịch, có lỗi thì hoàn tác.')

page('Thuật toán thanh toán và tính tiền')
picture('flow_payment','Lưu đồ thu tiền và hoàn tất đơn',12.4)
p('Tổng đơn bằng tổng của số lượng nhân đơn giá từng món. Tiền thừa bằng tiền khách đưa trừ tổng đơn. Điểm cộng bằng phần nguyên của tổng đơn chia 10.000, chỉ cộng cho khách có đăng ký. Nếu quyền, trạng thái hoặc phiên bản đơn không còn hợp lệ ở bước ghi, chương trình báo lỗi và hoàn tác giao dịch.')

page('Thiết kế giao diện và thao tác')
p('Giao diện dùng nền kem, thẻ trắng và xanh đậm cho thao tác chính. Vàng biểu thị đang phục vụ, đỏ được dùng cho xoá hoặc huỷ. Logo và góc bo tạo nhận diện thống nhất, nhưng control vẫn là ttk gốc để giữ thao tác bằng bàn phím và trạng thái vô hiệu hoá.')
picture('staff','Không gian phục vụ của tài khoản nhân viên')
p('Nhân viên có ba mục điều hướng: đơn hàng, sơ đồ bàn và thực đơn. Màn hình bắt đầu là sơ đồ bàn để giảm số thao tác khi có khách mới. Trạng thái luôn có chữ đi kèm màu. Mỗi bàn đưa ra hành động gọi món hoặc tiếp tục đơn phù hợp với tình trạng hiện tại.')
p('Quản lý có thêm tổng quan, khách hàng, nhân viên, tài khoản và báo cáo doanh thu. Những màn hình nhiều dữ liệu có bộ lọc hoặc tìm kiếm. Các nút có phản hồi khi rê chuột, nhận focus và bị vô hiệu hoá để người sử dụng nhận biết thao tác đang được phép.')

page('CHƯƠNG 3. CÀI ĐẶT GIẢI PHÁP',1)
h('Môi trường cài đặt và chạy')
table('Môi trường đã xác nhận khi kiểm tra',['Thành phần','Phiên bản / cấu hình'],[
['Hệ điều hành','Windows 11, bản dựng 26300.'],['Python','3.14.3.'],['Tk','8.6.15, dùng qua tkinter/ttk.'],['pyodbc','5.3.0.'],['SQL Server','Express 64-bit, phiên bản 16.0.1200.5.'],['Kết nối mặc định',r'THINKPADX1\SQLSEVER, Windows Authentication.'],['Database mặc định','QuanLyQuanCoffee.'],['Thư viện cài thêm','requirements.txt chỉ khai báo pyodbc>=5.2,<6.']], [5,10.5])
p('Tkinter, csv, decimal, datetime và unittest thuộc thư viện chuẩn Python. Hình PNG được Tkinter đọc trực tiếp, nên ứng dụng không cần Pillow để chạy. Thiết lập tạo dữ liệu mẫu chỉ khi bảng tương ứng còn rỗng, vì vậy chạy lại không nhân đôi tài khoản hoặc thực đơn.')
h('Quy trình khởi động')
p('Người dùng chuẩn bị SQL Server và driver ODBC, cài thư viện rồi chạy database_setup.py. Sau đó chạy main.py để mở đăng nhập. CoffeeApp quản lý việc chuyển giữa LoginWindow và MainWindow. Thông tin kết nối có thể thay đổi bằng COFFEE_SQL_SERVER và COFFEE_SQL_DATABASE khi chuyển sang máy khác.')
p('Ứng dụng không yêu cầu một IDE cụ thể. Các tệp Python có thể mở bằng công cụ lập trình của học phần hoặc một trình soạn thảo hỗ trợ Python. Quy trình chạy và kiểm thử dùng trực tiếp Python tại thư mục dự án.')

page('Đối chiếu nội dung học phần: Bài 1 đến Bài 4')
table('Vận dụng bốn nội dung đầu của học phần',['Nội dung','Lý thuyết / thực hành','Minh chứng trong mã'],[
['Bài 1. Giới thiệu về lập trình Python và các công cụ lập trình','3 / 3','main.py là điểm chạy. requirements.txt khai báo phụ thuộc. Chạy chương trình và kiểm thử bằng Python.'],['Bài 2. Các khái niệm cơ bản','3 / 3','Biến, kiểu dữ liệu, phép toán, import, lớp và đối tượng. Dùng int, str, bool và Decimal.'],['Bài 3. Các biểu thức điều kiện và vòng lặp','6 / 6','if kiểm tra giá, quyền và trạng thái. for gộp món, tính tổng và ghi từng dòng CSV.'],['Bài 4. Hàm trong Python','6 / 6','price, phone, report_period, transaction và các phương thức CoffeeService tách nhiệm vụ xử lý.']], [5.7,2.6,7.2])
h('Đoạn mã gộp món và kiểm tra số lượng')
p('Trích models/service.py, bên trong save_order. Đầu vào cart là các cặp mã sản phẩm và số lượng. Hàm integer kiểm tra kiểu và miền giá trị. Hai dòng cùng mã được gộp trước khi ghi xuống dữ liệu.')
code('quantities = {}\nfor product_id, quantity in cart:\n    product_id = integer(product_id, "Mã sản phẩm", 1)\n    quantity = integer(quantity, "Số lượng", 1, 999)\n    quantities[product_id] = integer(\n        quantities.get(product_id, 0) + quantity,\n        "Số lượng mỗi sản phẩm", 1, 999)')
p('Ví dụ giỏ chứa (1, 2) và (1, 1), quantities trở thành {1: 3}. Nếu tổng vượt 999 hoặc số lượng là 1,5, hàm báo lỗi. Cách xử lý này thể hiện đồng thời biến, dict, vòng lặp, lời gọi hàm và điều kiện kiểm tra.')

page('Đối chiếu nội dung học phần: Bài 5 đến Bài 8')
table('Vận dụng bốn nội dung còn lại',['Nội dung','Lý thuyết / thực hành','Minh chứng trong mã'],[
['Bài 5. Xử lý chuỗi','3 / 3','strip loại khoảng trắng. Chuẩn hoá tìm kiếm không dấu. Kiểm tra số điện thoại và định dạng số tiền.'],['Bài 6. Xử lý danh sách','3 / 3','Danh sách bản ghi, list comprehension, lọc và sắp xếp dữ liệu bảng. Giỏ và ánh xạ dùng thêm dict.'],['Bài 7. Xử lý tập tin','3 / 3','open theo context manager, encoding UTF-8 BOM, csv.writer ghi báo cáo. Nhật ký lỗi ghi có giới hạn.'],['Bài 8. Thiết kế giao diện với tkinter','3 / 3','Tk, Toplevel, ttk.Frame, Entry, Button, Combobox, Treeview, StringVar, pack và grid.']], [5.7,2.6,7.2])
p('Tổng thời lượng tám nội dung là 30 tiết lý thuyết và 30 tiết thực hành. Các kiến thức được kết hợp trong cùng quy trình xử lý, từ nhận dữ liệu trên giao diện đến kiểm tra, tính toán và lưu kết quả.')
h('Đoạn mã ghi các dòng sản phẩm ra tệp')
p('Trích ui/statistics_tab.py. Các dòng thông tin kỳ và tiêu đề đã được ghi trước vòng lặp này. products là list các dict do dịch vụ thống kê trả về.')
code('for row in products:\n    name = row["TenSP"]\n    if name.lstrip().startswith(("=", "+", "-", "@")):\n        name = "\'" + name\n    writer.writerow([row["MaSP"], name,\n                     row["SoLuong"], str(row["DoanhThu"])])')
p('Đoạn mã thể hiện xử lý chuỗi, duyệt danh sách, truy cập dict và ghi tập tin. Nội dung tiếng Việt được giữ nhờ cách mở tệp với encoding="utf-8-sig". Đường dẫn lưu được chọn bằng hộp thoại, không ghi mặc định vào một thư mục cố định.')

page('Giao dịch dữ liệu và bảo vệ lịch sử')
h('Đoạn mã quản lý giao dịch')
p('Trích db_connection.py. Context manager cho phép các phương thức nghiệp vụ dùng một kết nối trong một khối xử lý. Nếu khối hoàn thành thì commit, có ngoại lệ thì rollback, sau cùng luôn đóng kết nối.')
code('@contextmanager\ndef transaction():\n    connection = connect()\n    try:\n        connection.execute("SET XACT_ABORT ON")\n        yield connection.cursor()\n        connection.commit()\n    except Exception:\n        connection.rollback()\n        raise\n    finally:\n        connection.close()')
p('Khi thanh toán, trạng thái đơn, điểm khách và trạng thái bàn phải cùng thành công. Nếu chỉ ghi được một phần, ứng dụng có thể báo bàn trống trong khi đơn chưa hoàn tất. Gom các thay đổi trong giao dịch giúp giữ dữ liệu nhất quán. Cơ chế commit và rollback của pyodbc được đối chiếu với tài liệu của thư viện [9].')
h('Các biện pháp kiểm soát phù hợp quy mô')
p('Giá trị người dùng nhập được truyền bằng tham số truy vấn. Các tên bảng động chỉ chọn từ danh sách cấu hình cố định. Kiểm tra quyền được thực hiện lại trong CoffeeService. Mật khẩu dùng PBKDF2-HMAC-SHA256 với salt ngẫu nhiên và phép so sánh phù hợp, dựa trên các hàm trong hashlib và hmac [10].')
p('PhienBan chống ghi đè khi hai cửa sổ cùng sửa một đơn. Khoá giao dịch cho nhóm đơn hàng giúp điều phối các thao tác liên quan đến bàn và thanh toán. Chương trình ưu tiên tính đúng và sự rõ ràng của nghiệp vụ, thay vì thêm nhiều tầng kiến trúc cho số lượng chức năng nhỏ.')

page('Kết quả giao diện: đăng nhập và tổng quan')
picture('login','Màn hình đăng nhập',13.3)
picture('dashboard','Màn hình tổng quan của quản lý')
p('Đăng nhập tập trung vào tài khoản và mật khẩu. Sau khi xác thực, quản lý thấy các chỉ số hoạt động và biểu đồ, còn nhân viên đi thẳng tới sơ đồ bàn. Các ảnh minh hoạ được chụp khi chạy ứng dụng với dữ liệu trình diễn trong cơ sở dữ liệu kiểm tra riêng.')

page('Kết quả giao diện: thực đơn và gọi món')
picture('products','Danh mục thực đơn với tìm kiếm và các thao tác quản lý')
picture('order','Gọi món với hai cà phê đen và một cà phê sữa')
p('Màn hình gọi món đặt thực đơn bên trái và giỏ bên phải. Tổng tiền hiển thị gần nút lưu để dễ kiểm tra trước khi kết thúc thao tác. Ví dụ minh hoạ có 2 × 25.000 + 1 × 30.000 = 80.000 đồng. Các nút điều chỉnh chỉ hoạt động khi chọn dòng trong giỏ.')

page('Kết quả giao diện: thu tiền và báo cáo')
picture('payment','Hộp thoại thu tiền hiển thị tiền thừa 20.000 đồng',7.2)
picture('statistics','Báo cáo 7 ngày với doanh thu 1.200.000 đồng và 8 hoá đơn')
p('Với tiền khách đưa 100.000 đồng và tổng đơn 80.000 đồng, chương trình hiển thị tiền thừa 20.000 đồng. Bộ dữ liệu trình diễn của báo cáo có 8 đơn đã thanh toán, doanh thu 1.200.000 đồng và trung bình 150.000 đồng mỗi đơn. Đây là số liệu kiểm tra, không phải doanh thu kinh doanh thực tế.')

page('CHƯƠNG 4. ĐÁNH GIÁ VÀ KIỂM THỬ',1)
h('Mục tiêu và phương pháp kiểm thử')
p('Kiểm thử tập trung vào tính đúng của dữ liệu, quyền sử dụng và các luồng thao tác chính. Các trường hợp hợp lệ kiểm tra kết quả mong muốn. Các trường hợp sai kiểm tra việc từ chối và bảo đảm không để dữ liệu dở dang. Những tình huống đồng thời và thay đổi dữ liệu giữa hai cửa sổ được kiểm tra để bảo vệ bàn, thanh toán và lịch sử.')
p('Bộ tests/test_coffee.py sử dụng unittest, SQL Server thật và Tkinter. Mỗi lần chạy tạo một cơ sở dữ liệu CoffeeTest_ với tên ngẫu nhiên, thiết lập dữ liệu rồi tự xoá sau kiểm tra. Dữ liệu vận hành của quán không được dùng làm dữ liệu thử. Một số hộp thoại thông báo được thay thế trong kiểm thử để không phải xác nhận thủ công, nhưng các phép tính và giao dịch vẫn chạy thật.')
h('Kết quả lần kiểm tra ngày 04/10/2026')
p('Kết quả chạy lại là 19 kiểm thử đạt trong 19,801 giây. Không ghi nhận trường hợp thất bại trong lần chạy này. Kết quả này xác nhận các tình huống đã được kiểm tra trong môi trường nêu ở Chương 3, không thay thế việc kiểm tra trên một máy có cấu hình SQL Server khác.')
code('Ran 19 tests in 19.801s\n\nOK')
h('Dữ liệu thử bổ sung cho ảnh và xuất tệp')
p('Cơ sở dữ liệu trình diễn có 8 bàn, 5 món khởi đầu, một khách hàng Nguyễn Thanh An và tài khoản nhân viên Trần Ngọc Mai. Tạo 7 đơn theo 7 ngày và một đơn 80.000 đồng để kiểm tra thu tiền. Sau khi thanh toán, số liệu báo cáo là 30 cà phê đen, 15 cà phê sữa và tổng doanh thu 1.200.000 đồng.')
p('Các tên và số điện thoại trong ảnh dùng để trình diễn thao tác. Sau khi lấy bằng chứng, cơ sở dữ liệu trình diễn đã được xoá. Tệp CSV được mở lại để kiểm tra mã hoá, tiêu đề, tổng tiền và số đơn.')

page('Kết quả kiểm thử: dữ liệu và quyền')
table('Các trường hợp kiểm thử tự động, nhóm thứ nhất',['STT','Trường hợp và kết quả mong muốn','Kết quả'],[
['1','Băm mật khẩu khác salt, đăng nhập đúng, từ chối sai mật khẩu và chuỗi chèn SQL.','Đạt'],['2','Chạy thiết lập lại không nhân đôi dữ liệu mẫu.','Đạt'],['3','Từ chối giá sai, số điện thoại sai, tên rỗng và điểm không nguyên.','Đạt'],['4','Thêm, đọc, sửa, xoá các danh mục theo quyền.','Đạt'],['5','Tạo, sửa, chuyển bàn và xoá đơn chưa thanh toán theo quy tắc.','Đạt'],['6','Lỗi trong lưu đơn hoàn tác thay đổi.','Đạt'],['7','Thanh toán cộng điểm và giữ lịch sử, không sửa/xoá đơn đã thanh toán.','Đạt'],['8','Huỷ hoặc xoá đơn cập nhật đúng trạng thái bàn.','Đạt'],['9','Nhân viên bị từ chối thao tác quản trị ở tầng nghiệp vụ.','Đạt'],['10','Không tự xoá, tự khoá hoặc hạ quyền quản lý đang dùng.','Đạt']], [1.3,12.5,1.7])
p('Các kết quả được đối chiếu với dữ liệu SQL sau thao tác, không chỉ dựa trên thông báo giao diện. Ví dụ, khi lưu đơn gặp món không hợp lệ, kiểm tra cả đơn và trạng thái bàn để chắc chắn không còn cập nhật một phần.')

page('Kết quả kiểm thử: lịch sử, đồng thời và giao diện')
table('Các trường hợp kiểm thử tự động, nhóm thứ hai',['STT','Trường hợp và kết quả mong muốn','Kết quả'],[
['11','Sửa giá menu không đổi giá đã lưu. Món ngừng bán bị hạn chế đúng.','Đạt'],['12','Bản đơn cũ không ghi đè hoặc thanh toán sau khi đơn đã đổi.','Đạt'],['13','Hai thao tác giữ cùng bàn chỉ cho một đơn đang xử lý thành công.','Đạt'],['14','Hai thao tác thanh toán chỉ hoàn tất đơn và cộng điểm một lần.','Đạt'],['15','Lỗi SQL khi cộng điểm hoàn tác cả thanh toán và trạng thái bàn.','Đạt'],['16','Báo cáo đúng ngày, tháng, năm nhuận, không tính đơn huỷ/chưa thanh toán.','Đạt'],['17','CRUD tài khoản, đổi mật khẩu và kiểm tra phiên bị khoá.','Đạt'],['18','Đăng nhập, mở các trang, biểu mẫu và luồng gọi món bằng Tkinter.','Đạt'],['19','Nút, combobox, spinbox bo góc hoạt động bằng chuột, bàn phím và đúng trạng thái disabled.','Đạt']], [1.3,12.5,1.7])
p('Hai thử nghiệm đồng thời kiểm tra kết quả ghi xuống SQL Server. Kiểm thử phiên bản bảo đảm một cửa sổ cũ phải tải lại trước khi tiếp tục. Kiểm thử giao diện có cả kích thước cửa sổ nhỏ và tương tác control, giúp phát hiện trường hợp có nút đẹp nhưng không bấm được.')

page('Đối chiếu phép tính, xuất tệp và lưu trữ')
table('Các kiểm tra bổ sung khi chạy ứng dụng',['Mã','Dữ liệu và thao tác','Kết quả xác nhận'],[
['BS01','2 cà phê đen × 25.000 và 1 cà phê sữa × 30.000. Nhập tiền nhận 100.000.','Tổng 80.000, thừa 20.000, khách nhận 8 điểm.'],['BS02','Xuất báo cáo 28/09–04/10/2026 rồi đọc lại CSV.','Có BOM UTF-8, đúng tiêu đề, 8 đơn và 1.200.000 đồng.'],['BS03','Đăng nhập tài khoản nhân viên.','Chỉ có đơn hàng, sơ đồ bàn và thực đơn.'],['BS04','Đóng cửa sổ ứng dụng, đăng nhập và mở lại bằng kết nối mới.','Đơn vẫn đã thanh toán, tổng 80.000 và điểm khách vẫn là 8.']], [1.5,7.2,6.8])
p('Đối chiếu chi tiết CSV: cà phê đen có 30 phần, doanh thu 750.000 đồng. Cà phê sữa có 15 phần, doanh thu 450.000 đồng. Hai giá trị cộng lại bằng 1.200.000 đồng, khớp với thẻ doanh thu trên giao diện. Số hoá đơn là 8 nên doanh thu trung bình mỗi đơn là 150.000 đồng.')
h('Đánh giá kết quả và giới hạn')
p('Ứng dụng đáp ứng luồng chính của một quán nhỏ, có giao diện hai vai trò, lưu trữ cơ sở dữ liệu, xử lý tính tiền và xuất báo cáo. Mã được chia theo trách nhiệm, các công thức dùng phép toán trực tiếp và được kiểm tra bằng dữ liệu cụ thể. Các nguồn tham khảo giúp giải thích lựa chọn thư viện và mô hình tổ chức chương trình.')
p('Chưa kiểm tra việc tắt đột ngột hệ điều hành hoặc mất điện. Chưa có cơ chế sao lưu tự động, nhập CSV, quản lý nguyên liệu hay tích hợp thu tiền trực tuyến. Các chức năng này không được trình bày như đã hoàn thành. Khi chạy trên máy khác, cần thiết lập driver, dịch vụ SQL và quyền tài khoản Windows phù hợp.')

page('KẾT LUẬN',1)
p('Đề tài đã xây dựng ứng dụng Coffee cho quy trình quản lý và phục vụ quán cà phê nhỏ. Chương trình có đăng nhập, phân quyền, danh mục, sơ đồ bàn, gọi món, thanh toán, điểm khách hàng và báo cáo doanh thu. Dữ liệu được lưu trong SQL Server, báo cáo được xuất ra CSV để đối chiếu.')
p('Các kiến thức của tám nội dung học phần được vận dụng trong các tình huống cụ thể. Điều kiện và vòng lặp dùng kiểm tra và tổng hợp món. Hàm giúp tách nhiệm vụ. Chuỗi, list và dict dùng chuẩn hoá dữ liệu, giỏ hàng và kết quả truy vấn. Xử lý tập tin phục vụ xuất báo cáo. Tkinter/ttk tạo giao diện và liên kết thao tác với nghiệp vụ.')
p('Qua kiểm thử, các tình huống đã kiểm tra cho kết quả đúng về giá lịch sử, trạng thái bàn, thanh toán, điểm khách, báo cáo và quyền sử dụng. Những trường hợp lỗi có kiểm tra hoàn tác dữ liệu. Giao diện sử dụng nhãn rõ ràng, các vùng chức năng gọn và control thống nhất để hỗ trợ thao tác nhanh.')
p('Hướng phát triển phù hợp trước tiên là nhập danh mục từ CSV có kiểm tra từng dòng và bổ sung quy trình sao lưu. Sau đó có thể mở rộng quản lý nguyên liệu hoặc in phiếu khi có yêu cầu sử dụng thực tế. Những phần mở rộng cần được bổ sung theo nhu cầu, giữ cách tổ chức mã đơn giản và tránh làm phức tạp phạm vi học phần.')

page('TÀI LIỆU THAM KHẢO',1)
refs=[
('abdelbend, Restaurant Management System, kho mã nguồn GitHub.','https://github.com/abdelbend/Restaurant-Management-System'),
('Python Software Foundation, tkinter — Python interface to Tcl/Tk.','https://docs.python.org/3/library/tkinter.html'),
('mkleehammer và cộng sự, pyodbc — Python ODBC bridge.','https://github.com/mkleehammer/pyodbc'),
('Python Software Foundation, csv — CSV File Reading and Writing.','https://docs.python.org/3/library/csv.html'),
('Python Software Foundation, decimal — Decimal fixed-point and floating-point arithmetic.','https://docs.python.org/3/library/decimal.html'),
('Python Software Foundation, unittest — Unit testing framework.','https://docs.python.org/3/library/unittest.html'),
('abdelbend, db.py trong Restaurant Management System.','https://github.com/abdelbend/Restaurant-Management-System/blob/main/db.py'),
('Microsoft, rowversion (Transact-SQL).','https://learn.microsoft.com/en-us/sql/t-sql/data-types/rowversion-transact-sql'),
('pyodbc Wiki, Database Transaction Management.','https://github.com/mkleehammer/pyodbc/wiki/Database-Transaction-Management'),
('Python Software Foundation, hashlib — Secure hashes and message digests.','https://docs.python.org/3/library/hashlib.html')]
for i,(title,url) in enumerate(refs,1):
    pp=p(f'[{i}] {title} Truy cập ngày 04/10/2026.')
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    hyp=OxmlElement('w:hyperlink')
    hyp.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r')
    rp=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts')
    f.set(qn('w:ascii'),'Times New Roman')
    f.set(qn('w:hAnsi'),'Times New Roman')
    rp.append(f)
    size=OxmlElement('w:sz')
    size.set(qn('w:val'),'26')
    rp.append(size)
    rr.append(rp)
    rr.append(OxmlElement('w:br'))
    tt=OxmlElement('w:t')
    tt.set(qn('xml:space'),'preserve')
    tt.text=url.replace('/', '/\u200b').replace('-', '-\u200b')
    rr.append(tt)
    hyp.append(rr)
    pp._p.append(hyp)
p('[11] Trường Đại học Nam Cần Thơ, mẫu Báo cáo học phần Lập trình Python - Thực hành, tệp BaoCao_LapTrinhPython_Maubaocao.pdf, tài liệu học phần.')
p('[12] Trường Đại học Nam Cần Thơ, Hướng dẫn thực hành Lập trình với Python, Module 1–2 và Module 3, tài liệu học phần.')

page('PHỤ LỤC. CÀI ĐẶT VÀ SỬ DỤNG',1)
h('Chuẩn bị và chạy chương trình')
p('Cài Python có Tkinter, SQL Server và ODBC Driver 17 for SQL Server hoặc driver SQL Server tương thích đã được chương trình hỗ trợ. Kiểm tra dịch vụ SQL Server đang chạy. Tài khoản Windows dùng chạy Python phải có quyền truy cập cơ sở dữ liệu. Bước thiết lập lần đầu cần quyền tạo database.')
p('Mở cửa sổ dòng lệnh tại thư mục dự án và thực hiện theo thứ tự:')
code('python -m pip install -r requirements.txt\npython database_setup.py\npython main.py')
p('Nếu máy khác có tên server khác, đặt biến môi trường trước khi chạy. Ví dụ trong PowerShell, thay tên server phù hợp và giữ cùng cấu hình cho cả bước thiết lập và khởi động:')
code('$env:COFFEE_SQL_SERVER = "TENMAY\\TENCUMSQL"\n$env:COFFEE_SQL_DATABASE = "QuanLyQuanCoffee"')
p('Lần thiết lập đầu tạo tài khoản admin với mật khẩu Admin@123. Sau đăng nhập, dùng Đổi mật khẩu. Chạy thiết lập lại giữ nguyên dữ liệu đã có. Không xoá hoặc đổi tên database khi ứng dụng đang sử dụng.')
h('Chạy bộ kiểm thử')
code('python -B -m unittest discover -s tests -v')
p('Bộ kiểm thử cần quyền tạo và xoá database thử. Không dùng tên database vận hành làm tên thử. Tuỳ chọn -B tránh tạo tệp bytecode. Tài nguyên logo và control trong assets phải được giữ khi sao chép dự án.')

page('Hướng dẫn thao tác và kịch bản trình bày')
h('Thao tác của quản lý và nhân viên')
p('Quản lý đăng nhập, kiểm tra tổng quan, chuẩn bị thực đơn và bàn, tạo nhân viên rồi cấp tài khoản theo vai trò. Khi thêm dữ liệu, điền các trường bắt buộc và lưu. Tìm kiếm dùng tên hoặc nội dung hiển thị. Khi bản ghi đã được tham chiếu, dùng ngừng bán hoặc cập nhật phù hợp thay vì cố xoá lịch sử.')
p('Nhân viên chọn bàn trống, gọi món, chọn khách nếu tích điểm và lưu đơn. Để thêm món, mở Tiếp tục đơn. Khi thu tiền, kiểm tra tổng, nhập tiền đã nhận rồi xác nhận hoàn tất. Nếu khách chưa thanh toán, chọn quay lại và giữ đơn đang xử lý.')
p('Quản lý mở Báo cáo doanh thu, chọn ngày hoặc tháng, xem dữ liệu rồi Xuất CSV. Tệp được lưu ở vị trí người dùng chọn. Nút F5 làm mới trang. Nếu ứng dụng báo dữ liệu đã đổi ở cửa sổ khác, đóng và mở lại đơn để lấy phiên bản mới.')
h('Kịch bản trình bày gắn với minh chứng')
table('Thứ tự trình bày chức năng',['Bước','Nội dung cần thể hiện'],[
['1','Đăng nhập quản lý, giới thiệu tổng quan và cách bố trí giao diện.'],['2','Thêm hoặc sửa một món, tìm kiếm và kiểm tra thông tin trong SQL Server.'],['3','Đăng nhập nhân viên, chỉ ra ba mục chức năng và giới hạn quyền.'],['4','Tạo đơn 2 cà phê đen và 1 cà phê sữa, giải thích tổng 80.000 đồng.'],['5','Nhập 100.000 đồng, đối chiếu thừa 20.000, thanh toán và kiểm tra bàn.'],['6','Mở báo cáo, xuất CSV và đối chiếu tổng với dữ liệu hiển thị.'],['7','Giải thích transaction, dict giỏ hàng và chỉ ra kết quả kiểm thử.']], [1.5,14])
p('Khi giải thích mã, bắt đầu từ main.py, chuyển sang hàm nghiệp vụ rồi tới transaction. Khi giải thích lỗi, chọn một ví dụ cụ thể như số lượng bằng 0 hoặc đơn đã thay đổi. Không cần đọc hết mã nguồn theo thứ tự để mô tả một thao tác.')

update=OxmlElement('w:updateFields')
update.set(qn('w:val'),'true')
doc.settings._element.append(update)
out=ROOT/'docs/BaoCao_QuanLyQuanCaPhe.docx'
doc.save(out)
print(out)
