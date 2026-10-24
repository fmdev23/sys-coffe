"""Kiểm tra cấu trúc, đối chiếu nguồn và lập ma trận chứng cứ."""
import sys,ast,json,re,hashlib,textwrap,unicodedata
from pathlib import Path
from zipfile import ZipFile
from datetime import datetime,timedelta,timezone
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'verification'
NOW=datetime.now(timezone(timedelta(hours=7))).isoformat(timespec='seconds')
SOURCE=ROOT/'docs/BaoCao_QuanLyQuanCaPhe.docx'
candidate=ROOT/'docs/BaoCao_QuanLyQuanCaPhe_DRAFT.docx'
if not candidate.exists():candidate=ROOT/'docs/BaoCao_QuanLyQuanCaPhe_FINAL.docx'
sha=lambda value:hashlib.sha256(value).hexdigest()
load=lambda name:json.loads((OUT/name).read_text(encoding='utf-8-sig'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with ZipFile(SOURCE) as z:oldmedia={sha(z.read(n)):n for n in z.namelist() if n.startswith('word/media/')}; oldxml=E.fromstring(z.read('word/document.xml'))
with ZipFile(candidate) as z:
    assert z.testzip() is None
    xml=E.fromstring(z.read('word/document.xml'));styles=E.fromstring(z.read('word/styles.xml'))
    media={sha(z.read(n)):n for n in z.namelist() if n.startswith('word/media/')}
    all_parts=z.namelist()
text='\n'.join(xml.xpath('//w:t/text()',namespaces=ns))
doc=Document(candidate)
captions=[''.join(p._p.xpath('.//w:t/text()')) for p in doc.paragraphs if p.style.name=='Caption']
assert len(doc.tables)==15 and len(doc.inline_shapes)==13 and oldmedia.keys()==media.keys()
assert ';' not in text
assert not re.search(r'\[TODO\]|\bCN[0-9]+\b|\bG\.[0-9]+\b|\ufffd|[\u200b\u200c\u200d\ufeff]',text)
assert all(item in text for item in ['Lê Minh Phú','236717','DH23TIN02','Bùi Thị Diễm Trinh','Visual Studio Code','19.656s','22/10/2026'])
assert all(item not in text for item in ['18,003','16,156','guarded','contextmanager','CoffeeService.login'])
groups={}
for cap in captions:
    match=re.match(r'(Hình|Bảng) (\d+)(?:\.(\d+))? ',cap);assert match,cap
    kind,ch,num=match.groups();key=(kind,ch if num else 'Giới thiệu')
    groups.setdefault(key,[]).append(int(num or ch))
for key,values in groups.items():assert values==list(range(1,len(values)+1)),(key,values)
rows=[]
for i,table in enumerate(doc.tables,1):
    borders=table._tbl.xpath('./w:tblPr/w:tblBorders')[0]
    assert {E.QName(c).localname for c in borders}=={'top','left','bottom','right','insideH','insideV'}
    assert all(b.get('{'+ns['w']+'}val')=='single' for b in borders)
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    for row in table.rows:assert row._tr.xpath('./w:trPr/w:cantSplit')
    rows.append({'table':i,'rows':len(table.rows),'columns':len(table.columns),'full_grid':True,'no_split_xml':True,'repeat_header':True})
for section in doc.sections:
    assert abs(section.page_width.cm-21)<.01 and abs(section.page_height.cm-29.7)<.01
    for attr,target in [('left_margin',3.5),('right_margin',2),('top_margin',2),('bottom_margin',2),('header_distance',1),('footer_distance',1)]:
        assert abs(getattr(section,attr).cm-target)<.01,(attr,getattr(section,attr).cm)
normal=doc.styles['Normal'];assert normal.font.name=='Times New Roman' and normal.font.size.pt==13
assert abs(normal.paragraph_format.line_spacing-1.2)<.001
assert abs(normal.paragraph_format.first_line_indent.cm-1)<.01
assert normal.paragraph_format.space_before.pt==6 and normal.paragraph_format.space_after.pt==0
fields=xml.xpath('//w:instrText/text()|//w:fldSimple/@w:instr',namespaces=ns)
assert sum(' TOC ' in f for f in fields)==3
assert any('TOC \\o "1-2" \\h' in f for f in fields)
assert sum('SEQ Figure' in f for f in fields)==12 and sum('SEQ Table' in f for f in fields)==15
links=xml.xpath('//w:hyperlink[@w:anchor]',namespaces=ns);assert len(links)==60
bookmark_names=set(xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns))
assert all(link.get('{'+ns['w']+'}anchor') in bookmark_names for link in links)
word=load('truong_word.json');assert word['fields_match'] and word['no_split_rows']
assert len(word['page_references'])==60
old_dims=oldxml.xpath('//wp:inline/wp:extent',namespaces=ns)
new_dims=xml.xpath('//wp:inline/wp:extent',namespaces=ns)
assert len(old_dims)==len(new_dims)==13
assert all((a.get('cx'),a.get('cy'))==(b.get('cx'),b.get('cy')) for a,b in zip(old_dims,new_dims))
source_trees={p:ast.parse((ROOT/p).read_text(encoding='utf-8-sig')) for p in [r['path'] for r in load('ma_nguon.json')]}
class SqlWhitespace(ast.NodeTransformer):
    def visit_Constant(self,node):
        if isinstance(node.value,str) and re.match(r'\s*(UPDATE|SELECT|INSERT|DELETE|SET)\b',node.value):
            node.value=re.sub(r'\s+',' ',node.value).strip()
        return node
def comparable(node):
    import copy
    return ast.dump(SqlWhitespace().visit(copy.deepcopy(node)))
code_checks=[]
for p in doc.paragraphs:
    if p.style.name=='Code' and p.text.lstrip().startswith('def '):
        tree=ast.parse(textwrap.dedent(p.text))
        for node in tree.body:
            if isinstance(node,ast.FunctionDef):
                matches=[path for path,source in source_trees.items() for original in ast.walk(source) if isinstance(original,ast.FunctionDef) and comparable(node)==comparable(original)]
                assert matches,node.name
                code_checks.append({'function':node.name,'files':matches,'ast_equivalent_except_sql_whitespace':True})
assert len(code_checks)==2
pdf=PdfReader(OUT/'bao_cao_kiem_tra.pdf');assert len(pdf.pages)==34
assert 'Error! Reference source not found.' not in '\n'.join(page.extract_text() for page in pdf.pages)
changed=[]
for f in load('dau_vao_sha256.json')['files']:
    p=ROOT/f['path']
    if not p.exists() or sha(p.read_bytes())!=f['sha256']:changed.append(f['path'])
assert not changed,changed
sources=load('ma_nguon.json')
syntax=[]
for row in sources:
    p=ROOT/row['path'];ast.parse(p.read_text(encoding='utf-8-sig'));syntax.append({'path':row['path'],'syntax':'ĐÚNG','sha256':sha(p.read_bytes())})
source_check={'checked_at':NOW,'files':syntax,'count':len(syntax),'no_bytecode_generated':True}
(OUT/'kiem_tra_cu_phap.json').write_text(json.dumps(source_check,ensure_ascii=False,indent=2),encoding='utf-8')
result={'checked_at':NOW,'candidate':str(candidate),'zip_valid':True,'paragraphs':len(doc.paragraphs),'tables':rows,'media_count':13,'media_bytes_identical':True,'image_dimensions_unchanged':True,'captions':captions,'caption_groups':{str(k):v for k,v in groups.items()},'internal_links':len(links),'page_references_match':True,'table_rows_not_split_word':True,'code_snippets':code_checks,'a4_margins_body_style':True,'source_and_input_files_unchanged':True,'input_files_checked':len(load('dau_vao_sha256.json')['files']),'pages':len(pdf.pages),'visual_review':'Lưu riêng tại xem_trang.json sau khi xem đủ ảnh trang.'}
(OUT/'kiem_tra_docx.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['tables','captions','caption_groups','code_snippets']},ensure_ascii=False,indent=2))
