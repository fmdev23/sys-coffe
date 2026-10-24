"""Ghi nhận đầu vào và trích cấu trúc để kiểm định, không sửa dự án."""
import ast
import hashlib
import json
import re
import subprocess
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'verification'
NOW = datetime.now(timezone(timedelta(hours=7))).isoformat(timespec='seconds')

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def hide(value):
    for secret in ('Admin@123', 'NhanVien@123', 'Abcd@123'):
        value = value.replace(secret, '(mật khẩu mẫu đã ẩn)')
    return value

git = {}
for label, args in [('status', ['status', '--short', '--branch']), ('branch', ['branch', '--show-current']),
                    ('commit', ['rev-parse', 'HEAD']), ('remote', ['remote', '-v'])]:
    result = subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True, encoding='utf-8')
    git[label] = {'exit_code': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr}
(OUT / 'git_ban_dau.json').write_text(json.dumps({'time': NOW, **git}, ensure_ascii=False, indent=2), encoding='utf-8')

paths = []
for folder in ['models', 'ui', 'tests', 'assets', 'docs', 'sodo', '.report_work']:
    paths.extend(p for p in (ROOT / folder).rglob('*') if p.is_file())
paths.extend(p for p in ROOT.iterdir() if p.is_file())
snapshot = [{'path': p.relative_to(ROOT).as_posix(), 'bytes': p.stat().st_size, 'sha256': sha(p)} for p in sorted(set(paths))]
(OUT / 'dau_vao_sha256.json').write_text(json.dumps({'time': NOW, 'files': snapshot}, ensure_ascii=False, indent=2), encoding='utf-8')

for filename, target in [('BaoCao_LapTrinhPython_Maubaocao.pdf', 'mau_python.txt'),
                         ('Quyet-dinh-ban-hanh-Qui-dinh-trinh-bay-luan-van-(DNC).pdf', 'quy_dinh_dnc.txt')]:
    reader = PdfReader(ROOT / 'docs' / filename)
    text = '\n\n'.join(f'TRANG {i}\n{p.extract_text()}' for i, p in enumerate(reader.pages, 1))
    (OUT / target).write_text(text, encoding='utf-8')
    print(filename, len(reader.pages), 'trang')

path = ROOT / 'docs/BaoCao_QuanLyQuanCaPhe.docx'
document = Document(path)
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(path) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    parts = [{'part': n, 'sha256': hashlib.sha256(z.read(n)).hexdigest(), 'bytes': len(z.read(n))} for n in z.namelist()]
    media = [n for n in z.namelist() if n.startswith('word/media/')]
    fields = xml.xpath('//w:instrText/text()', namespaces=ns)
    captions = []
    for p in xml.xpath('//w:body/w:p', namespaces=ns):
        if p.xpath('./w:pPr/w:pStyle[@w:val="Caption"]', namespaces=ns):
            captions.append(''.join(p.xpath('.//w:t/text()', namespaces=ns)))
    sections = [etree.tostring(s, encoding='unicode') for s in xml.xpath('//w:sectPr', namespaces=ns)]

paragraphs = [{'index': i, 'style': p.style.name, 'text': hide(p.text)} for i, p in enumerate(document.paragraphs)]
tables = [[[hide(c.text) for c in r.cells] for r in t.rows] for t in document.tables]
data = {'time': NOW, 'source_sha256': sha(path), 'paragraphs': paragraphs, 'tables': tables,
        'captions': captions, 'fields': fields, 'media': media, 'sections': sections, 'parts': parts,
        'internal_links': len(xml.xpath('//w:hyperlink[@w:anchor]', namespaces=ns)),
        'all_text': hide('\n'.join(xml.xpath('//w:t/text()', namespaces=ns)))}
(OUT / 'bao_cao_dau_vao.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

source = []
for folder in [ROOT, ROOT/'models', ROOT/'ui', ROOT/'tests']:
    for p in sorted(folder.glob('*.py')):
        tree = ast.parse(p.read_text(encoding='utf-8-sig'), filename=str(p))
        functions = []
        for n in ast.walk(tree):
            if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)):
                functions.append({'name': n.name, 'line': n.lineno, 'end_line': n.end_lineno})
        source.append({'path': p.relative_to(ROOT).as_posix(), 'sha256': sha(p), 'functions': functions})
(OUT / 'ma_nguon.json').write_text(json.dumps(source, ensure_ascii=False, indent=2), encoding='utf-8')
print('AST:', len(source), 'tệp, hợp lệ')
print('Báo cáo:', len(paragraphs), 'đoạn,', len(tables), 'bảng,', len(media), 'ảnh,', len(captions), 'chú thích')
